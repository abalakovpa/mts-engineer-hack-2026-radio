#!/usr/bin/env python3
"""Cross-check generated artifacts and source integrity; do not certify KPI."""
from pathlib import Path
from collections import Counter
import hashlib
import csv
import json
import math
import re
import zipfile

from openpyxl import load_workbook
from pptx import Presentation
from calculate import capacity

ROOT=Path(__file__).resolve().parents[2]
HASHES={
    'Кабель LCF78-50JA-A0.pdf':'b341944deccf6ebebfae2dda0c94b13c62b7794b52d848265328377c858ac739',
    'Параметры планировки и затуханий.png':'aee4bd1e84094f74f7c38f02bddccb99785cf87a77fb1047cb1fbf392bf9e593',
    'Перечень вспомогательного оборудования.xlsx':'b92688685fc219b924e9cf9a4fa47147c13ae18747d3cc75691412e7d1b39cb0',
    'Планировка первого этажа.dwg':'c5f7b901edcdb073f1248ab7dc169d6b85c4b08f80b6cef5e137ef67635f00d7',
    'Планировка первого этажа.png':'13760675a4151a114537d31cbef25631eb8de775d5e3f5289adb034db6891c96',
    'Кабель LCF12-50J.pdf':'047726c01e5ee745efc214c7a5b47d274fb9d2f5e7e984a25bb4bcb75b7fc6bd',
}

DOCX_HASHES={'docs/final/task.docx': 'e1c72d32e324dd875d3fc007eb3e72a341e2e75649f6c0a1561d76d69e4bdf85', 'docs/online/calculation-results.docx': 'c44c557685fa45b2cf0460f31376831d0e9a57bf4de379eee10cf453d735da43', 'docs/online/design-approach.docx': '5f305327a7af2631f225e8f41dc5617955de8b13ef1284af441c52c286421e2e', 'docs/online/equipment-selection.docx': 'f6ce420aab0eb7b7a891c34007ac493b6e7892a8260f7c88efa32cc6c146abb7', 'docs/online/methodology.docx': '65282c4dbc96dd8393e1c2e72796148559d5dc9fdd40852ba6c6655158b25af2'}

def main():
    for path in ROOT.rglob('*.md'):
        if '.git' in path.parts or '.venv' in path.parts:
            continue
        for target in re.findall(r'\]\(([^\n)]+)\)',path.read_text()):
            target=target.strip().strip('<>')
            if re.match(r'^[a-zA-Z][a-zA-Z0-9+.-]*:',target) or target.startswith('#'):
                continue
            local=target.split('#',1)[0]
            assert (path.parent/local).exists(),(str(path.relative_to(ROOT)),target)
    for name,digest in DOCX_HASHES.items():
        assert hashlib.sha256((ROOT/name).read_bytes()).hexdigest()==digest,name
    for name,digest in HASHES.items():
        actual=hashlib.sha256((ROOT/'docs/final/source-materials'/name).read_bytes()).hexdigest()
        assert actual==digest,name
    capacity()  # Independent S/S-1 and exact arithmetic checks.
    data=json.loads((ROOT/'design/final/candidate.json').read_text())
    ants=data['antennas'];count=len(ants);assert count in (32,48) and len({a['id'] for a in ants})==count
    per_zone=count//4;subgroups=per_zone//4
    assert Counter(a['zone'] for a in ants)=={z:per_zone for z in range(1,5)}
    for z in range(1,5):
        if data['topology'].get('zone_modes',{}).get(str(z))=='unequal':
            assert Counter(a['subgroup'] for a in ants if a['zone']==z)=={0:2,1:2,2:4}
            # Ideal power fractions at the zone input sum to unity.
            fractions={'direct':.25,'pair':.125,'quartet':.0625}
            assert sum(fractions[a['power_class']] for a in ants if a['zone']==z)==1
            for a in [a for a in ants if a['zone']==z]:
                delta={'direct':-10*math.log10(2)-.5,'pair':0,'quartet':10*math.log10(2)}[a['power_class']]
                assert math.isclose(a['distribution_adjustment_db'],delta,abs_tol=1e-10)
        else:assert all(sum(a['zone']==z and a['subgroup']==g for a in ants)==4 for g in range(1,subgroups+1))
    routes=data['trunks']+data['subfeeds']+data['drops']
    expected_routes=4+4*subgroups+count
    assert len(routes)==expected_routes
    assert len({t['id'] for t in routes})==expected_routes
    for r in routes:
        measured=sum(math.dist(a,b) for a,b in zip(r['points'],r['points'][1:]))
        assert math.isclose(measured,r['length'],abs_tol=1e-8),(r['id'],measured,r['length'])
        assert r['cable']=='jumper' or r['length']>0,(r['id'],'Zero-length feeder')
        if r['cable']=='jumper':assert r['jumper_length_m']==3 and r['length']<1e-6
    assert data['local_jumper_count']==2*sum(r['cable']=='jumper' for r in routes)
    for cable,key in [('7/8','cable78_installed'),('1/2','cable12_installed')]:
        total=2*sum(r['length'] for r in routes if r['cable']==cable)
        assert math.isclose(total,data[key],abs_tol=1e-8)
    points=data['rf_points'];assert len(points)==count*12
    assert len({(r['antenna'],r['frequency_mhz'],r['d_horizontal_m']) for r in points})==count*12
    for a in ants:
        for f in (1800,2100,2600):
            vals=[r['rsrp_dbm'] for r in points if r['antenna']==a['id'] and r['frequency_mhz']==f]
            assert len(vals)==4 and all(u>v for u,v in zip(vals,vals[1:]))
            selected=[r for r in points if r['antenna']==a['id'] and r['frequency_mhz']==f]
            alpha12={1800:10.09,2100:10.99,2600:11.84+(12.66-11.84)*2/3}[f]
            alpha78={1800:5.21,2100:5.67,2600:6.39}[f]
            loss=18.830899869919438+a.get('distribution_adjustment_db',0)+a['length78_m']*alpha78/100+a['length12_m']*alpha12/100+a['additional_jumper_loss_db']
            pilot=23-10*math.log10(2*{1800:100,2100:75,2600:50}[f])
            for r in selected:
                predicted=pilot-loss+3-(20*math.log10(f)-27.55+25*math.log10(math.hypot(r['d_horizontal_m'],5.5)))
                assert math.isclose(loss,r['total_loss_db'],abs_tol=1e-8)
                assert math.isclose(predicted,r['rsrp_dbm'],abs_tol=1e-8)
    wb=load_workbook(ROOT/'equipment/bill-of-materials.xlsx',data_only=True)
    assert wb['RF и RSRP'].max_row==count*12+1 and wb['Трассы'].max_row==expected_routes+1
    for row,r in zip(wb['RF и RSRP'].iter_rows(min_row=2,values_only=True),points):
        assert tuple(row[:3])==(r['antenna'],r['frequency_mhz'],r['d_horizontal_m'])
        assert math.isclose(row[-1],r['rsrp_dbm'],abs_tol=1e-8)
    spec=list(wb['Спецификация'].iter_rows(min_row=2,values_only=True))
    expected={'In06':count,'PC211':8,'PC221':6,'PC250':9 if count==32 else 1,
              'PC251':0 if count==32 else 8,'PC252':subgroups*8}
    for name,quantity in expected.items():
        matched=[r for r in spec if r[1].startswith(name)]
        if quantity:assert len(matched)==1 and matched[0][3]==quantity,(name,matched)
        else:assert not matched,(name,matched)
    for token,cable,key in [('1/2','1/2','cable12_installed'),('7/8','7/8','cable78_installed')]:
        feeder=next(r for r in spec if r[1].startswith('Фидер') and token in r[1])
        connector=next(r for r in spec if r[1].startswith('Разъем') and token in r[1])
        assert feeder[3]==math.ceil(data[key]*1.1)
        assert connector[3]==4*sum(r['cable']==cable for r in routes)
    jumper=next(r for r in spec if r[1].startswith('Джамп-кабель'))
    assert jumper[3]==45+2*count+data['local_jumper_count']
    for f in (ROOT/'equipment/bill-of-materials.xlsx',ROOT/'presentation/defense.pptx'):
        with zipfile.ZipFile(f) as z:assert z.testzip() is None
    worst=min(r['rsrp_dbm'] for r in points if r['frequency_mhz']==2600 and r['d_horizontal_m']==20)
    coverage_file=ROOT/'design/final/coverage/coverage-report.json'
    if count==32:assert coverage_file.exists(), 'Actual coverage report required for selected design'
    if coverage_file.exists():
        coverage=json.loads(coverage_file.read_text())
        assert coverage['antennas']==count
        assert math.isclose(coverage['located_reference_m2'],2051.56,abs_tol=1e-6)
        assert coverage['unknown_reference_m2']==45.52
        assert len(coverage['scenarios'])==39
        assert coverage['area_model_sha256']==hashlib.sha256((ROOT/'design/final/coverage/area-model.json').read_bytes()).hexdigest()
        model=json.loads((ROOT/'design/final/coverage/area-model.json').read_text())
        assert len(model['groups'])==36 and not model['metadata']['area_joins_are_rf_walls']
        assert all(j['length_m']<=.25+1e-8 for j in model['area_joins'])
        assert all(w['layer'] not in ('A-AREA','A-FLOR') for w in model['walls'])
        assert [g['rooms'] for g in model['groups'] if abs(g['relative_area_difference'])>.1]==[['114.2']]
        by_key={(r['scenario'],r['frequency_mhz']):r for r in coverage['scenarios']}
        for f in (1800,2100,2600):
            nominal=by_key[('P1_n2.5',f)]['reference_percent_2100']
            assert nominal<=2051.56/2100*100+1e-6
            for name in ('P1_n3','P1_n3_closed_gaps','P1_extra_RF_3dB','P1_closed_gaps','P1_two_ports_total','P2_carrier23','P1_cable78_plus5pct','P1_geometry_uncertain_fail',
                         'P1_closed_gaps_RF3','P1_geometry_uncertain_fail_RF3','P1_closed_gaps_geometry_uncertain_fail_RF3'):
                assert by_key[(name,f)]['reference_percent_2100']<=nominal+1e-6
            assert by_key[('P1_closed_gaps_RF3',f)]['reference_percent_2100']<=by_key[('P1_extra_RF_3dB',f)]['reference_percent_2100']+1e-6
            assert by_key[('P1_closed_gaps_geometry_uncertain_fail_RF3',f)]['reference_percent_2100']<=by_key[('P1_closed_gaps_RF3',f)]['reference_percent_2100']+1e-6
            assert by_key[('P1_n2',f)]['reference_percent_2100']>=nominal-1e-6
            for name in ('P1_n3','P1_n3_closed_gaps','P1_extra_RF_3dB','P1_closed_gaps_geometry_uncertain_fail_RF3'):
                assert by_key[(name,f)]['reference_percent_2100']>=95,'Selected reserve objective not met'
        assert wb['Покрытие сценарии'].max_row==40
        assert coverage['design_fingerprint']==hashlib.sha256((ROOT/'design/final/candidate.json').read_bytes()).hexdigest()
        for row,result in zip(wb['Покрытие сценарии'].iter_rows(min_row=2,values_only=True),coverage['scenarios']):
            assert row[0]==result['scenario'] and row[1]==result['frequency_mhz']
            assert math.isclose(row[5],result['reference_percent_2100'],abs_tol=1e-8)
            assert result['passed_95']==(result['reference_percent_2100']>=95)
        import numpy as np
        with np.load(ROOT/'design/final/coverage/validated-fields.npz') as field:
            assert len(field['points'])==coverage['receiver_cells']
            assert math.isclose(float(field['weights'].sum()),2051.56,abs_tol=1e-6)
            for f in (1800,2100,2600):
                predicted=float(field['weights'][field[f'rsrp_{f}']>-90].sum()/2100*100)
                assert math.isclose(predicted,by_key[('P1_n2.5',f)]['reference_percent_2100'],abs_tol=1e-8)
        resolution=json.loads((ROOT/'design/final/coverage/resolution-0.5.json').read_text())
        assert resolution['design_fingerprint']==coverage['design_fingerprint']
        assert resolution['area_model_sha256']==coverage['area_model_sha256']
        control=json.loads((ROOT/'design/final/coverage/geometry-control.json').read_text())
        assert control['design_fingerprint']==coverage['design_fingerprint']
        assert control['area_model_sha256']==coverage['area_model_sha256']
        assert control['area_join_max_m']==.15 and control['step_m']==.5
        assert all(control['scores'][s]>=95 for s in ('n3','reserve3','closed_reserve3'))
    quality_path=ROOT/'design/final/quality/quality-report.json'
    if quality_path.exists():
        quality=json.loads(quality_path.read_text())
        assert quality['design_fingerprint']==coverage['design_fingerprint']
        assert quality['area_model_sha256']==coverage['area_model_sha256']
        assert not quality['full_kpi_compliance']
        for key,path in [('lte_fields_sha256','design/final/coverage/validated-fields.npz'),
                         ('dcs_fields_sha256','design/final/verification/dcs-field.npz')]:
            assert quality[key]==hashlib.sha256((ROOT/path).read_bytes()).hexdigest()
        params={s['id']:s for s in quality['scenarios']}
        assert len(quality['lte_results'])==33 and len(quality['dcs_results'])==12
        with np.load(ROOT/'design/final/coverage/validated-fields.npz') as field, np.load(ROOT/'design/final/quality/quality-fields.npz') as saved:
            weights=field['weights']
            assert np.array_equal(saved['points'],field['points']) and np.array_equal(saved['weights'],weights)
            for result in quality['lte_results']:
                p=params[result['scenario']];f=result['frequency_mhz'];prb={1800:100,2100:75,2600:50}[f]
                level=field[f'rsrp_n3_{f}'] if p['field']=='P1_n3' else field[f'rsrp_{f}']
                if p['field']=='P2':level=level-10*math.log10(6)
                signal=10**(level/10)
                thermal=10**((-174+10*math.log10(15000)+p['nf_db'])/10)
                interference=0 if p['interference_dbm'] is None else 10**(p['interference_dbm']/10)
                j=thermal+interference
                predicted_sinr=10*np.log10(p['beta']*signal/j)
                rssi=prb*(4*signal+8*2*p['rho']*p['beta']*signal+12*j)
                predicted_rsrq=10*np.log10(prb*signal/rssi)
                for name,t in [('sinr_ge17_percent',17),('sinr_ge7_percent',7),('sinr_ge1_percent',1)]:
                    assert math.isclose(result[name],float(weights[predicted_sinr>=t].sum()/21),abs_tol=1e-8)
                assert math.isclose(result['rsrq_gt_minus12_percent'],float(weights[predicted_rsrq>-12].sum()/21),abs_tol=1e-8)
                assert math.isclose(result['rsrp_gt_minus90_percent'],float(weights[level>-90].sum()/21),abs_tol=1e-8)
                if result['scenario']=='L4':assert result['rsrq_gt_minus12_percent']==0
                if result['scenario']==quality['base_scenario']:
                    assert np.allclose(saved[f'sinr_{f}'],predicted_sinr) and np.allclose(saved[f'rsrq_{f}'],predicted_rsrq)
            with np.load(ROOT/'design/final/verification/dcs-field.npz') as dcs:
                for result in quality['dcs_results']:
                    c=dcs['rx_dbm']+result['bcch_dbm']-40
                    ci=c-result['interference_dbm_200khz']
                    assert math.isclose(result['ci_gt20_percent'],float(weights[ci>20].sum()/21),abs_tol=1e-8)
                    assert math.isclose(result['rx_gt_minus85_percent'],float(weights[c>-85].sum()/21),abs_tol=1e-8)
                assert np.allclose(saved['dcs_ci'],dcs['rx_dbm']+70)
        assert wb['LTE качество'].max_row==34 and wb['Бюджет помех'].max_row==10 and wb['DCS C-I'].max_row==13
        for row,result in zip(list(wb['LTE качество'].values)[1:],quality['lte_results']):
            assert row[0].startswith(result['scenario']+' — ') and row[1]==result['frequency_mhz']
            assert math.isclose(row[9],result['rsrq_gt_minus12_percent'],abs_tol=1e-8)
            assert row[11]==result['including_rsrp_met']
        print('PASS: 33 conditional LTE quality scenarios; 12 DCS C/I scenarios; independent RSSI/SINR arithmetic')
    review_file=ROOT/'design/final/verification/engineering-review.json'
    if review_file.exists():
        review=json.loads(review_file.read_text())
        assert review['design_fingerprint']==coverage['design_fingerprint']
        assert review['manufactured_feeder_segments']==2*sum(r['cable']!='jumper' for r in routes)
        assert review['total_jumper_count']==45+2*count+data['local_jumper_count']
        assert review['unused_input_loads_mts_only']==16 and not review['full_kpi_compliance']
        ports=list(csv.DictReader((ROOT/'design/final/verification/port-plan.csv').open()))
        assert len(ports)==137 and len({r['target'] for r in ports})==137
        wired=[r['source'] for r in ports if r['source']!='Нагрузка 50 Ом']
        assert len(wired)==len(set(wired)), 'A device output port was used twice'
        antenna_ports={r['target'] for r in ports if r['target'].startswith('A')}
        assert antenna_ports=={f'{a["id"]}:P{m}' for a in ants for m in (1,2)}
        devices=list(csv.DictReader((ROOT/'design/final/verification/device-locations.csv').open()))
        assert len(devices)==len({r['id'] for r in devices})==71 and wb['Размещение'].max_row==72
        assert Counter(r['model'] for r in devices)=={k:v for k,v in expected.items() if v}
        with np.load(ROOT/'design/final/verification/dcs-field.npz') as dcs, np.load(ROOT/'design/final/coverage/validated-fields.npz') as lte:
            diff=40-10*math.log10(2)-.5-(23-10*math.log10(200))
            assert np.allclose(dcs['rx_dbm'],lte['rsrp_1800']+diff)
            share=float(dcs['weights'][dcs['rx_dbm']>-85].sum()/2100*100)
            assert math.isclose(share,review['dcs_scenarios'][0]['reference_percent_2100'],abs_tol=1e-8)
            assert all(r['reference_percent_2100']<=share+1e-8 for r in review['dcs_scenarios'])
        assert wb['Портовый план'].max_row==138 and wb['Комплектность'].max_row==11
        final_prs=Presentation(ROOT/'presentation/defense.pptx')
        assert len(final_prs.slides)==6
        template=ROOT/'presentation/source-template.pptx'
        if template.exists():
            template_hash='83f1fd8e22427bed6954db9cccb7145e5662e14aa89208947e42e5fab1598a28'
            assert hashlib.sha256(template.read_bytes()).hexdigest()==template_hash
            assert final_prs.slide_width==9144000 and final_prs.slide_height==5143500
            with zipfile.ZipFile(template) as original, zipfile.ZipFile(ROOT/'presentation/defense.pptx') as result:
                fonts={name:original.read(name) for name in original.namelist() if name.startswith('ppt/fonts/')}
                assert fonts and all(result.read(name)==content for name,content in fonts.items())
            notes='\n'.join(s.notes_slide.notes_text_frame.text for s in final_prs.slides)
            assert 'github' not in notes.lower() and '[заполнить]' not in notes
            import fitz
            with fitz.open(ROOT/'presentation/defense-preview.pdf') as rendered:
                assert len(rendered)==6
                rendered_text='\n'.join(page.get_text() for page in rendered)
                for value in ('13 UE','104','32 In06','96,37','1366','179'):
                    assert value in rendered_text,value
            status=json.loads((ROOT/'presentation/presentation-status.json').read_text())
            assert status['template_sha256']==template_hash and status['official_template_available']
            assert status['rendered_from_pptx']
            if quality_path.exists():
                assert status['quality_report_sha256']==hashlib.sha256(quality_path.read_bytes()).hexdigest()
                assert re.search(r'25\s*→\s*2\s*\(k=1,10\)\s*→\s*1',rendered_text)
                assert 'SINR' in rendered_text and 'RSRQ' in rendered_text
        for slide in final_prs.slides:
            assert slide.has_notes_slide
            for shape in slide.shapes:
                assert shape.left>=0 and shape.top>=0
                assert shape.left+shape.width<=final_prs.slide_width+1 and shape.top+shape.height<=final_prs.slide_height+1
        print('PASS: final port plan, 64 antenna ports, 71 device locations, DCS field, completeness sheet, final presentation')
    manifest_file=ROOT/'submission/manifest.json'
    if manifest_file.exists():
        manifest=json.loads(manifest_file.read_text())
        assert manifest['design_fingerprint']==coverage['design_fingerprint']
        if quality_path.exists():
            assert manifest['quality_report_sha256']==hashlib.sha256(quality_path.read_bytes()).hexdigest()
            assert manifest['conditional_quality_calculated']
        assert len(manifest['files']) in (4,5) and not manifest['full_kpi_compliance']
        assert manifest['presentation_included']==(len(manifest['files'])==5)
        expected_files={'01-capacity.pdf','02-floor-design.pdf','03-rsrp.pdf','04-equipment.xlsx'}
        if manifest['presentation_included']:expected_files.add('05-defense.pptx')
        assert {item['name'] for item in manifest['files']}==expected_files
        assert {p.name for p in (ROOT/'submission/final').iterdir()}==expected_files
        for item in manifest['files']:
            p=ROOT/'submission/final'/item['name']
            assert p.stat().st_size==item['bytes'] and hashlib.sha256(p.read_bytes()).hexdigest()==item['sha256']
        assert (ROOT/'submission/final/04-equipment.xlsx').read_bytes()==(ROOT/'equipment/bill-of-materials.xlsx').read_bytes()
        if manifest['presentation_included']:
            assert (ROOT/'submission/final/05-defense.pptx').read_bytes()==(ROOT/'presentation/defense.pptx').read_bytes()
        archive=ROOT/'submission'/manifest['archive']['name']
        assert archive.stat().st_size<=18_000_000 and hashlib.sha256(archive.read_bytes()).hexdigest()==manifest['archive']['sha256']
        with zipfile.ZipFile(archive) as z:
            assert z.testzip() is None and set(z.namelist())=={r['name'] for r in manifest['files']}
            for item in manifest['files']:assert hashlib.sha256(z.read(item['name'])).hexdigest()==item['sha256']
        print(f'PASS: {len(manifest["files"])} separate submission artifacts; ZIP hashes and size')
    print(f'PASS: 11 source hashes; 7 capacity variants; {count} antennas; {expected_routes} routes ×2; {count*12} RF points; XLSX quantities; 6 slides')
    print(f'RF scenario at 20 m on weakest B7 branch: {worst:.2f} dBm; KPI readiness NOT established')

if __name__=='__main__':main()
