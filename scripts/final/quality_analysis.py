#!/usr/bin/env python3
"""Conditional DL quality and interference budgets using the existing CAD fields.

No interferer sites, EARFCN/ARFCN or measured receiver parameters are invented.
Unknown interference is represented by an explicit aggregate receiver level.
"""
from pathlib import Path
import hashlib
import json
import math

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from shapely.geometry import shape
from build_project import style

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'design/final/quality'
BASE='L2'
FREQUENCIES=(1800,2100,2600)
SCENARIOS=[
    dict(id='L0',label='Только тепловой шум',beta=.5,rho=.8,nf_db=7,interference_dbm=None,field='P1'),
    dict(id='L1',label='Помехи −125 дБм/RE',beta=.5,rho=.8,nf_db=7,interference_dbm=-125,field='P1'),
    dict(id='L2',label='Помехи −115 дБм/RE',beta=.5,rho=.8,nf_db=7,interference_dbm=-115,field='P1'),
    dict(id='L3',label='Помехи −105 дБм/RE',beta=.5,rho=.8,nf_db=7,interference_dbm=-105,field='P1'),
    dict(id='L4',label='Равная EPRE данных и CRS',beta=1,rho=.8,nf_db=7,interference_dbm=-115,field='P1'),
    dict(id='L5',label='Полная загрузка',beta=.5,rho=1,nf_db=7,interference_dbm=-115,field='P1'),
    dict(id='L6',label='n=3',beta=.5,rho=.8,nf_db=7,interference_dbm=-115,field='P1_n3'),
    dict(id='L7',label='Пилот P2',beta=.5,rho=.8,nf_db=7,interference_dbm=-115,field='P2'),
    dict(id='L8',label='NF=9 дБ',beta=.5,rho=.8,nf_db=9,interference_dbm=-115,field='P1'),
    dict(id='L9',label='Помехи −95 дБм/RE',beta=.5,rho=.8,nf_db=7,interference_dbm=-95,field='P1'),
    dict(id='L10',label='Помехи −85 дБм/RE',beta=.5,rho=.8,nf_db=7,interference_dbm=-85,field='P1'),
]


def mw(dbm):return np.power(10.,np.asarray(dbm)/10)
def db(value):return 10*np.log10(value)
def noise_dbm(bandwidth_hz,nf_db):return -174+10*math.log10(bandwidth_hz)+nf_db


def lte_metrics(signal_dbm,beta,rho,nf_db,interference_dbm):
    """Two TX ports, CRS-bearing symbol: four CRS RE and eight data RE/RB.

    beta is PDSCH EPRE on each port / CRS EPRE on that port. rho is data
    occupancy in the measurement symbols. Streams are ideally separated;
    this per-stream SINR surrogate does not predict a UE receiver's RI/CQI.
    """
    assert beta>0 and 0<=rho<=1
    signal=mw(signal_dbm)
    disturbance=mw(noise_dbm(15000,nf_db))
    if interference_dbm is not None:disturbance+=mw(interference_dbm)
    sinr=db(beta*signal/disturbance)
    rsrq=db(signal/((4+16*rho*beta)*signal+12*disturbance))
    return sinr,rsrq,float(db(disturbance))


def area_share(values,threshold,weights,strict=False):
    passed=values>threshold if strict else values>=threshold
    return float(weights[passed].sum()/2100*100)


def area_quantile(signal,weights,fraction):
    """Lowest level needed to cover fraction of the complete 2100 m² floor."""
    required=2100*fraction
    if required>weights.sum():return None
    order=np.argsort(-signal,kind='stable')
    index=np.searchsorted(np.cumsum(weights[order]),required,side='left')
    return float(signal[order[index]])


def interference_budget(signal,weights,beta=.5,rho=.8,nf_db=7):
    constraints={f'SINR_{threshold}_{int(fraction*100)}':area_quantile(signal,weights,fraction)+10*math.log10(beta)-threshold
                 for threshold,fraction in ((17,.4),(7,.8),(1,.95))}
    coefficient=10**1.2-(4+16*rho*beta)
    if coefficient<=0:
        return dict(feasible_rsrq=False,rsrq_ceiling_db=-10*math.log10(4+16*rho*beta),constraints_dbm=constraints,
                    max_disturbance_dbm=None,max_external_interference_dbm=None)
    ratio=12/coefficient
    constraints['RSRQ_-12_80']=area_quantile(signal,weights,.8)-10*math.log10(ratio)
    limit=min(constraints.values())
    noise=noise_dbm(15000,nf_db)
    external=float(db(mw(limit)-mw(noise))) if limit>noise else None
    # The RSSI criterion is strict; operational examples stay 0.1 dB below this boundary.
    return dict(feasible_rsrq=True,rsrq_ceiling_db=-10*math.log10(4+16*rho*beta),
                constraints_dbm=constraints,max_disturbance_dbm=limit,
                max_external_interference_dbm=external,thermal_floor_dbm=noise)


def plot_field(points,values,model,name,title,label,vmin,vmax):
    fig,ax=plt.subplots(figsize=(9,14))
    chart=ax.scatter(points[:,0],points[:,1],c=values,s=7,marker='s',cmap='RdYlGn',vmin=vmin,vmax=vmax)
    for group in model['groups']:ax.plot(*shape(group['geometry']).exterior.xy,color='#333333',lw=.35)
    ax.set_aspect('equal');ax.set_xlim(-2,39);ax.set_ylim(-2,66)
    ax.set_xlabel('CAD X, м');ax.set_ylabel('CAD Y, м');ax.set_title(title,fontsize=10)
    fig.colorbar(chart,ax=ax,label=label)
    fig.savefig(OUT/name,dpi=170,bbox_inches='tight');plt.close(fig)


def validate_equations():
    # Known two-port RSSI limits and monotonic changes in interference/power.
    signal=np.array([-85.,-95.,-105.])
    a,b,_=lte_metrics(signal,.5,.8,7,-115)
    c,d,_=lte_metrics(signal,.5,.8,7,-105)
    assert np.all(c<a) and np.all(d<b)
    _,equal,_=lte_metrics(np.array([0.]),1,.8,7,None)
    assert math.isclose(float(equal[0]),-10*math.log10(16.8),abs_tol=1e-10)
    assert np.all(equal<-12)
    _,full,_=lte_metrics(np.array([0.]),.5,1,7,None)
    assert math.isclose(float(full[0]),-10*math.log10(12),abs_tol=1e-10)
    weights=np.array([500.,800.,751.56]);budget=interference_budget(signal,weights)
    # Direct linear RSSI calculation agrees with the normalized form.
    for s,snr,q in zip(signal,a,b):
        power=10**(s/10);j=10**(noise_dbm(15000,7)/10)+10**(-115/10)
        rssi=100*((4+16*.8*.5)*power+12*j)
        assert math.isclose(q,10*math.log10(100*power/rssi),abs_tol=1e-12)
        assert math.isclose(snr,10*math.log10(.5*power/j),abs_tol=1e-12)
    # Use the budget's total disturbance, split into noise and interference.
    ext=budget['max_external_interference_dbm']
    if ext is not None:
        snr,q,_=lte_metrics(signal,.5,.8,7,ext-.1)
        assert all(area_share(snr,t,weights)>=p for t,p in ((17,40),(7,80),(1,95)))
        assert area_share(q,-12,weights,True)>=80


def main():
    validate_equations();OUT.mkdir(exist_ok=True)
    candidate=ROOT/'design/final/candidate.json';area=ROOT/'design/final/coverage/area-model.json'
    cov=json.loads((ROOT/'design/final/coverage/coverage-report.json').read_text())
    model=json.loads(area.read_text())
    assert cov['design_fingerprint']==hashlib.sha256(candidate.read_bytes()).hexdigest()
    assert cov['area_model_sha256']==hashlib.sha256(area.read_bytes()).hexdigest()
    with np.load(ROOT/'design/final/coverage/validated-fields.npz') as z:fields={key:z[key] for key in z.files}
    points=fields['points'];weights=fields['weights'];rows=[];budgets=[];saved={}
    for frequency in FREQUENCIES:
        nominal=fields[f'rsrp_{frequency}']
        signals={'P1':nominal,'P1_n3':fields[f'rsrp_n3_{frequency}'],'P2':nominal-10*math.log10(6)}
        for scenario in SCENARIOS:
            signal=signals[scenario['field']]
            sinr,rsrq,j=lte_metrics(signal,scenario['beta'],scenario['rho'],scenario['nf_db'],scenario['interference_dbm'])
            row=dict(scenario=scenario['id'],frequency_mhz=frequency,
                     sinr_ge17_percent=area_share(sinr,17,weights),sinr_ge7_percent=area_share(sinr,7,weights),
                     sinr_ge1_percent=area_share(sinr,1,weights),rsrq_gt_minus12_percent=area_share(rsrq,-12,weights,True),
                     rsrp_gt_minus90_percent=area_share(signal,-90,weights,True),disturbance_dbm_per_re=j)
            row['quality_targets_met']=bool(row['sinr_ge17_percent']>=40 and row['sinr_ge7_percent']>=80
                                           and row['sinr_ge1_percent']>=95 and row['rsrq_gt_minus12_percent']>=80)
            row['including_rsrp_met']=bool(row['quality_targets_met'] and row['rsrp_gt_minus90_percent']>=95)
            rows.append(row)
            if scenario['id']==BASE:
                saved[f'sinr_{frequency}']=sinr;saved[f'rsrq_{frequency}']=rsrq
                caption=f'{frequency} МГц, P1; L2: NF=7 дБ, β=0,5, ρ=0,8\nI = −115 дБм / 15 кГц; неизвестная площадь не засчитывается'
                plot_field(points,sinr,model,f'sinr-{frequency}.png',caption,'Сценарный SINR, дБ',0,35)
                plot_field(points,rsrq,model,f'rsrq-{frequency}.png',caption,'Сценарный RSRQ, дБ',-20,-9)
        for name,signal in signals.items():
            budget=interference_budget(signal,weights)
            budget.update(field=name,frequency_mhz=frequency,beta=.5,rho=.8,nf_db=7)
            if budget['max_external_interference_dbm'] is not None:
                sinr,rsrq,_=lte_metrics(signal,.5,.8,7,budget['max_external_interference_dbm']-.1)
                assert all(area_share(sinr,t,weights)>=p for t,p in ((17,40),(7,80),(1,95)))
                assert area_share(rsrq,-12,weights,True)>=80
            budgets.append(budget)
    with np.load(ROOT/'design/final/verification/dcs-field.npz') as z:
        assert np.array_equal(points,z['points']) and np.array_equal(weights,z['weights'])
        nominal_dcs=z['rx_dbm']
    dcs_rows=[];dcs_budgets=[]
    for power in (40.,40-10*math.log10(2),30.):
        carrier=nominal_dcs+power-40
        q95=area_quantile(carrier,weights,.95)
        dcs_budgets.append(dict(bcch_dbm=power,c95_dbm=q95,max_interference_dbm=q95-20))
        for interference in (-90.,-80.,-70.,-60.):
            ci=carrier-interference;passed=area_share(ci,20,weights,True)
            row=dict(bcch_dbm=power,interference_dbm_200khz=interference,ci_gt20_percent=passed,
                     rx_gt_minus85_percent=area_share(carrier,-85,weights,True))
            row['both_targets_met']=bool(passed>=95 and row['rx_gt_minus85_percent']>=95);dcs_rows.append(row)
            if power==40 and interference==-70:
                saved['dcs_ci']=ci
                plot_field(points,ci,model,'dcs-ci.png','DCS-1800: BCCH=40 дБм; I = −70 дБм / 200 кГц\nСценарный C/I >20 дБ; ARFCN и внешние источники не заданы','Сценарный C/I, дБ',0,40)
    report=dict(design_fingerprint=cov['design_fingerprint'],area_model_sha256=cov['area_model_sha256'],
                lte_fields_sha256=hashlib.sha256((ROOT/'design/final/coverage/validated-fields.npz').read_bytes()).hexdigest(),
                dcs_fields_sha256=hashlib.sha256((ROOT/'design/final/verification/dcs-field.npz').read_bytes()).hexdigest(),
                denominator_m2=2100,located_m2=float(weights.sum()),unknown_m2=2100-float(weights.sum()),
                base_scenario=BASE,scenarios=SCENARIOS,lte_results=rows,lte_budgets=budgets,
                dcs_results=dcs_rows,dcs_budgets=dcs_budgets,full_kpi_compliance=False,
                interpretation='Сценарный DL, идеальное разделение потоков MIMO, сильнейший антенный путь; помехи заданы суммарным уровнем на входе UE, частотный план не синтезируется',
                measurement_basis='RSRQ: CRS-содержащие символы, два TX-порта, 4 CRS RE/RB и 8 потенциальных PDSCH RE/RB',
                noise_psd_dbm_hz=-174,noise_re_hz=15000,rsrq_ceiling_equal_epre_load08_db=-10*math.log10(16.8))
    (OUT/'quality-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,allow_nan=False)+'\n')
    np.savez_compressed(OUT/'quality-fields.npz',points=points,weights=weights,**saved)
    workbook=load_workbook(ROOT/'equipment/bill-of-materials.xlsx')
    for title in ('LTE качество','Бюджет помех','DCS C-I'):
        if title in workbook:del workbook[title]
    ws=workbook.create_sheet('LTE качество')
    ws.append(['Сценарий','МГц','NF, дБ','I, дБм/15 кГц','β','ρ','SINR ≥17, %','SINR ≥7, %','SINR ≥1, %','RSRQ >−12, %','RSRP >−90, %','RSRP, SINR и RSRQ выполнены'])
    params={s['id']:s for s in SCENARIOS}
    for r in rows:
        s=params[r['scenario']]
        ws.append([s['id']+' — '+s['label'],r['frequency_mhz'],s['nf_db'],s['interference_dbm'] if s['interference_dbm'] is not None else 'Только шум',s['beta'],s['rho'],
                   r['sinr_ge17_percent'],r['sinr_ge7_percent'],r['sinr_ge1_percent'],r['rsrq_gt_minus12_percent'],r['rsrp_gt_minus90_percent'],r['including_rsrp_met']])
    style(ws,{chr(65+i): (40 if i==0 else 20) for i in range(12)})
    ws=workbook.create_sheet('Бюджет помех');ws.append(['Поле RSRP','МГц','β','ρ','NF, дБ','J предельное, дБм/15 кГц','I предельное, дБм/15 кГц','Условие'])
    for b in budgets:ws.append([b['field'],b['frequency_mhz'],b['beta'],b['rho'],b['nf_db'],b['max_disturbance_dbm'],b['max_external_interference_dbm'],'Граница для совокупности SINR/RSRQ; рабочий уровень минимум на 0,1 дБ ниже. Параметры сценарные.'])
    style(ws,{chr(65+i):(65 if i==7 else 24) for i in range(8)})
    ws=workbook.create_sheet('DCS C-I');ws.append(['BCCH, дБм','I, дБм/200 кГц','C/I >20, %','RxLev >−85, %','Оба KPI в сценарии','Условие'])
    for r in dcs_rows:ws.append([r['bcch_dbm'],r['interference_dbm_200khz'],r['ci_gt20_percent'],r['rx_gt_minus85_percent'],r['both_targets_met'],'Суммарная соканальная помеха на входе UE; ARFCN и внешние БС не заданы.'])
    style(ws,{'A':20,'B':28,'C':24,'D':24,'E':28,'F':65})
    workbook.save(ROOT/'equipment/bill-of-materials.xlsx')
    for r in rows:
        if r['scenario'] in ('L2','L3','L4','L6'):print(r)
    print('LTE allowed interference',[(b['field'],b['frequency_mhz'],b['max_external_interference_dbm']) for b in budgets])
    print('DCS budgets',dcs_budgets)

if __name__=='__main__':main()
