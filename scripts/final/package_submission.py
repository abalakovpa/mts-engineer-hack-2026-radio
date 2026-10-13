#!/usr/bin/env python3
"""Export five separate deliverables; validate the ZIP, never upload it."""
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
import hashlib
import json
import math
import re
import shutil
import subprocess
import zipfile
from openpyxl import load_workbook
from pptx import Presentation
from final_engineering import table

ROOT=Path(__file__).resolve().parents[2]
FINAL=ROOT/'docs/final'
EDITABLE=FINAL/'deliverables'
OUTPUT=ROOT/'submission/final'
ARCHIVE=ROOT/'submission/mts-radio-final.zip'
NAMES=['01-capacity.pdf','02-floor-design.pdf','03-rsrp.pdf','04-equipment.xlsx','05-defense.pptx']

def body(path):
    text=path.read_text()
    text=re.sub(r'^# .*\n','',text,count=1)
    text=re.sub(r'```mermaid\n.*?```','Схема приведена отдельным листом в этом документе.',text,flags=re.S)
    # Local source links are meaningful in the repository, but not in a standalone PDF.
    text=re.sub(r'\[([^\]]+)\]\((?!https?://)([^)]+)\)',r'\1',text)
    return text.strip()

def image_page(title,path,width=None,height=None):
    size=f'width={width}mm' if width else f'height={height or 205}mm,width=\\linewidth,keepaspectratio'
    # Nonfloating graphics keep the heading and drawing on the same page.
    return '\n\\newpage\n\n## '+title+'\n\n\\noindent\\includegraphics['+size+']{'+str(path)+'}\n'

def export_pdf(item):
    name,title= item
    source=EDITABLE/(name+'.md');target=OUTPUT/(name+'.pdf')
    cmd=['pandoc',str(source),'--from=markdown+tex_math_single_backslash','--standalone','--pdf-engine=xelatex',
         '--resource-path='+str(ROOT),'-V','mainfont=DejaVu Sans','-V','monofont=DejaVu Sans Mono',
         '-V','fontsize=10pt','-V','geometry:margin=18mm','-V','papersize=a4',
         '-V','colorlinks=true','-V','linkcolor=blue','--include-in-header='+str(EDITABLE/'pdf-header.tex'),
         '-M','title='+title,'-M','author=Павел Абалаков',
         '-M','date=13 октября 2026','-o',str(target)]
    result=subprocess.run(cmd,cwd=ROOT,text=True,capture_output=True)
    (EDITABLE/(name+'.build.log')).write_text(result.stdout+result.stderr)
    if result.returncode:raise RuntimeError(f'{name}: '+result.stderr[-2500:])
    if 'Missing character' in result.stderr:raise RuntimeError(f'{name}: missing glyphs')
    return target

def main():
    EDITABLE.mkdir(parents=True,exist_ok=True);OUTPUT.mkdir(parents=True,exist_ok=True)
    unexpected=set(p.name for p in OUTPUT.iterdir())-set(NAMES)
    assert not unexpected,('Do not silently delete unexpected submission files',unexpected)
    d=json.loads((ROOT/'design/final/candidate.json').read_text())
    cov=json.loads((ROOT/'design/final/coverage/coverage-report.json').read_text())
    audit=json.loads((ROOT/'design/final/verification/engineering-review.json').read_text())
    status=json.loads((ROOT/'presentation/presentation-status.json').read_text())
    assert audit['design_fingerprint']==cov['design_fingerprint']==hashlib.sha256((ROOT/'design/final/candidate.json').read_bytes()).hexdigest()
    capacity=body(FINAL/'capacity-revision.md')
    # The historical correction is retained in its source, not repeated as narrative in the final document.
    capacity=re.sub(r'Поправка к первоначальной редакции:.*?\n\n','',capacity,flags=re.S)
    capacity+='\n\n## Проверка и итоговая конфигурация\n\n'
    capacity+='Проверены точная арифметика, семь сочетаний диапазонов, S/S−1 и Erlang B. '
    capacity+='Для проекта выбран один общий LTE-пул МТС с B3/B1/B7 и две независимые MIMO-ветви; DCS — один пул, два TRX. '
    capacity+='Количество антенн 32 не увеличивает ёмкость пула. Средняя модель нагрузки не подтверждает временной перцентиль и фактическую скорость.\n\n'
    capacity+='Источники: task.docx — новые вводные и приложение; четыре авторских документа онлайн-этапа — прежняя модель; '
    capacity+='capacity-review.md и scripts/final/calculate.py — независимая арифметическая проверка. '
    capacity+='Исходные DOCX сохранены без изменений.\n'
    (EDITABLE/'01-capacity.md').write_text(capacity)
    project=body(FINAL/'design-project.md')+'\n\n\\newpage\n\n'
    project+=body(FINAL/'engineering-review.md')+'\n\n\\newpage\n\n'
    project+=body(FINAL/'reserve-review.md')+'\n\n\\newpage\n\n'
    project+=body(FINAL/'coverage-review.md')
    for title,path,w,h in [
        ('План первого этажа: антенны, делители, трассы и ВК','design/final/verification/final-layout.pdf',None,220),
        ('Схема подключения передатчиков БС','design/final/verification/bs-connections.pdf',170,None),
        ('Схема соединения распределительного оборудования','design/final/verification/distribution-tree.pdf',170,None),
        ('Карта RSRP B3: P1, n=2,5','design/final/coverage/rsrp-1800-walls.png',None,215),
        ('Карта RSRP B1: P1, n=2,5','design/final/coverage/rsrp-2100-walls.png',None,215),
        ('Карта RSRP B7: P1, n=2,5','design/final/coverage/rsrp-2600-walls.png',None,215),
        ('B7: дополнительные 3 дБ потерь','design/final/coverage/rsrp-2600-reserve3.png',None,215),
        ('B7: проёмы, спорный контур и ещё 3 дБ','design/final/coverage/rsrp-2600-joint-reserve3.png',None,215),
        ('Карта DCS: условная мощность BCCH 40 дБм','design/final/verification/dcs-1800-conditional.png',None,215),
        ('Принадлежность проектной LTE-ячейке: без внешней сети','design/final/verification/project-cell-assignment.png',None,210),
        ('Модель площади и неопределённые контуры','design/final/coverage/area-model.png',None,215),
    ]:project+=image_page(title,ROOT/path,w,h)
    project+='\n\\newpage\n\n## Координаты и группы антенн\n\n'
    project+='CAD X/Y, метры; высота антенн 7 м — допущение. Z — группа; g=0 — непосредственная, g=1 — парная, g=2 — четверная. P1/P2 обслуживаются раздельными деревьями. Номер — исходная привязка/целевая зона, не подтверждение принадлежности новой позиции отдельной комнате.\n\n'
    project+=table(['ID','Привязка','Z/g','X, м','Y, м','7/8, м','1/2, м'],
                   [(a['id'],a['room'],f'{a["zone"]}/{a["subgroup"]}',f'{a["x"]:.3f}',f'{a["y"]:.3f}',
                     f'{a["length78_m"]:.3f}',f'{a["length12_m"]:.3f}') for a in d['antennas']])
    project+='\n\n## Маршруты одной ветви Mj\n\nКаждый ID повторяется на M1 и M2. Геометрическая длина без закупочного запаса. '
    project+='Для связей с нулевым геометрическим расстоянием указан физический джампер 3 м.\n\n'
    project+=table(['ID','Соединение','Геометрия, м','Физическая длина на Mj, м'],
                   [(r['id'],r['cable'],f'{r["length"]:.3f}', '3 (заводской джампер)' if r['cable']=='jumper' else f'{r["length"]:.3f}')
                    for r in d['trunks']+d['subfeeds']+d['drops']])
    project+='\n\n## Координаты распределительных узлов\n\nНа каждой координате S/Q находятся отдельные устройства M1 и M2. '
    project+='Геометрические точки не задают размеры шкафов или точную монтажную раскладку.\n\n'
    project+=table(['Узел','Модель / количество','X, м','Y, м'],
                   [('ВК','6 PC221 + 8 PC211 + 1 PC250',f'{d["head"][0]:.3f}',f'{d["head"][1]:.3f}')]+
                   [(f'S-Z{z}','2 PC252',f'{d["trunks"][z-1]["points"][-1][0]:.3f}',f'{d["trunks"][z-1]["points"][-1][1]:.3f}') for z in range(1,5)]+
                   [(f'Q-Z{z}-{g}','2 '+('PC250' if g==1 else 'PC252'),f'{r["points"][-1][0]:.3f}',f'{r["points"][-1][1]:.3f}')
                    for z in range(1,5) for g in (1,2) for r in d['subfeeds'] if r['id']==f'F{z}-{g}'])
    project+='\n\n## Портовые обозначения\n\nH-band-Mj — гибрид PC221; TP-Zz-Mj — триплексер PC211; S-Zz-Mj — PC252; '
    project+='Q-Zz-1-Mj — PC250; Q-Zz-2-Mj — PC252; S-DCS — отдельный PC250. Полная таблица портов, включая входы БС, резерв операторов '
    project+='и оконечные антенны, включена в лист «Портовый план» отдельного артефакта 04-equipment.xlsx. '
    project+='Портовая маркировка I/O является проектной; точная маркировка корпуса устройства не предоставлена.\n'
    (EDITABLE/'02-floor-design.md').write_text(project)
    radio=body(FINAL/'rsrp-calculation.md')
    radio+='\n\n\\newpage\n\n## Проверяемый пример полного тракта\n\n'
    worst_id=max(d['rf_points'],key=lambda r:r['total_loss_db'])['antenna']
    a=next(a for a in d['antennas'] if a['id']==worst_id)
    radio+=f'{worst_id}: 7/8 — {a["length78_m"]:.6f} м; 1/2 — {a["length12_m"]:.6f} м; класс {a["power_class"]}. '
    radio+=f'Фиксированные потери {d["fixed_loss_db"]+a.get("distribution_adjustment_db",0):.8f} дБ; дополнительные местные джамперы {a["additional_jumper_loss_db"]:.2f} дБ. Последовательность: '
    radio+='мощность CRS RE → RF-потери → усиление 3 dBi → пространственная потеря → RSRP. '
    radio+='Один путь на одном порту, n=2,5, без препятствий.\n\n'
    for f,prb in ((1800,100),(2100,75),(2600,50)):
        rows=[r for r in d['rf_points'] if r['antenna']==worst_id and r['frequency_mhz']==f]
        pilot=23-10*math.log10(2*prb)
        radio+=f'### {f} МГц\n\nP_CRS_RE = {pilot:.6f} дБм; L_cable = {rows[0]["cable_loss_db"]:.6f} дБ; '
        radio+=f'L_RF = {rows[0]["total_loss_db"]:.6f} дБ.\n\n'
        radio+=table(['d по горизонтали, м','r, м','L_n, дБ','RSRP, дБм'],
                     [(r['d_horizontal_m'],f'{math.hypot(r["d_horizontal_m"],5.5):.4f}',
                       f'{20*math.log10(f)-27.55+25*math.log10(math.hypot(r["d_horizontal_m"],5.5)):.4f}',
                       f'{r["rsrp_dbm"]:.4f}') for r in rows])+'\n\n'
    radio+='## Препятствия и варианты определения мощности\n\nДля пересечённой перегородки вычитается 15 дБ, '
    radio+='для внешней стены или ЖБ колонны — 30 дБ. Число пересечений определяется лучом на карте, а не расстоянием само по себе. '
    radio+='P2 уменьшает каждую строку P1 на 7,7815 дБ; суммарные 23 дБм двух портов — на 3,0103 дБ. '
    radio+='Эти сценарии проверены отдельно; ни один не объявлен установленной настройкой БС.\n'
    radio+='\n\\newpage\n\n## Полная таблица 384 контрольных значений\n\n'
    radio+='P1, n=2,5; горизонтальные расстояния 5/10/15/20 м; высоты 7/1,5 м; без пересечений стен. '
    radio+='Номинально одинаково для двух независимых MIMO-путей. Кабели и компоненты включены в L_RF. '
    radio+='Это таблица уровней отдельных путей, не процент покрытия площади.\n\n'
    for f in (1800,2100,2600):
        radio+=f'### {f} МГц\n\n'
        rows=[]
        for a in d['antennas']:
            vals=sorted((r for r in d['rf_points'] if r['antenna']==a['id'] and r['frequency_mhz']==f),key=lambda r:r['d_horizontal_m'])
            rows.append([a['id'],f'{vals[0]["total_loss_db"]:.3f}']+[f'{r["rsrp_dbm"]:.2f}' for r in vals])
        radio+=table(['ID','L_RF, дБ','5 м','10 м','15 м','20 м'],rows)+'\n\n'
    radio+='\n\\newpage\n\n'+body(FINAL/'coverage-review.md')
    radio+='\n\n## Проверка выгрузки\n\n384 значения совпадают с листом «RF и RSRP» отдельной XLSX-спецификации. '
    radio+=f'{len(cov["scenarios"])} сценариев покрытия рассчитаны на актуальных координатах и потерях. Исходные паспорта и DOCX не изменены. '
    radio+='RSRQ, SINR, C/I и UL не выводятся из этих значений автоматически.\n'
    (EDITABLE/'03-rsrp.md').write_text(radio)
    jobs=[('01-capacity','Скорректированный расчёт количества секторов'),
          ('02-floor-design','Проектное решение первого этажа'),('03-rsrp','Расчёт RSRP: исходные данные и результаты')]
    with ThreadPoolExecutor(max_workers=3) as pool:list(pool.map(export_pdf,jobs))
    shutil.copyfile(ROOT/'equipment/bill-of-materials-draft.xlsx',OUTPUT/NAMES[3])
    shutil.copyfile(ROOT/'presentation/defense.pptx',OUTPUT/NAMES[4])
    # Page count and Unicode text extraction check exported files, not just the sources.
    files=[]
    for name in NAMES:
        p=OUTPUT/name;entry=dict(name=name,bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest())
        if p.suffix=='.pdf':
            info=subprocess.run(['pdfinfo',str(p)],capture_output=True,text=True,check=True).stdout
            entry['pages']=int(re.search(r'^Pages:\s+(\d+)',info,re.M).group(1))
            text=subprocess.run(['pdftotext','-layout',str(p),'-'],capture_output=True,text=True,check=True).stdout
            assert 'Павел Абалаков' in text and '\ufffd' not in text and len(text)>1500,name
            if name=='03-rsrp.pdf':assert worst_id in text and '384' in text,name
        else:
            with zipfile.ZipFile(p) as z:assert z.testzip() is None
        files.append(entry)
    wb=load_workbook(OUTPUT/NAMES[3],data_only=True)
    assert wb['RF и RSRP'].max_row==385 and wb['Портовый план'].max_row==138
    assert len(Presentation(OUTPUT/NAMES[4]).slides)==6
    with zipfile.ZipFile(ARCHIVE,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=9) as archive:
        for name in NAMES:
            entry=zipfile.ZipInfo(name,date_time=(2026,10,13,0,0,0));entry.compress_type=zipfile.ZIP_DEFLATED
            archive.writestr(entry,(OUTPUT/name).read_bytes(),compresslevel=9)
    with zipfile.ZipFile(ARCHIVE) as archive:
        assert archive.testzip() is None and archive.namelist()==NAMES
        for entry in files:assert hashlib.sha256(archive.read(entry['name'])).hexdigest()==entry['sha256']
    assert ARCHIVE.stat().st_size<=18_000_000,'ZIP exceeds conservative 18 MB limit'
    manifest=dict(date='2026-10-13',files=files,archive=dict(name=ARCHIVE.name,bytes=ARCHIVE.stat().st_size,
                  sha256=hashlib.sha256(ARCHIVE.read_bytes()).hexdigest()),
                  design_fingerprint=cov['design_fingerprint'],area_model_sha256=cov['area_model_sha256'],technical_export_checked=True,
                  full_kpi_compliance=False,uploaded_to_competition=False,
                  pending_profile_fields=status['pending_profile_fields'],official_template_available=False,
                  personal_rehearsal_done=False)
    (ROOT/'submission/manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(manifest,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
