#!/usr/bin/env python3
"""Final engineering supplements; all unknown radio inputs remain conditional."""
from pathlib import Path
import csv
import hashlib
import json
import math
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.collections import LineCollection
from matplotlib.patches import FancyBboxPatch
import numpy as np
from openpyxl import load_workbook
from shapely.geometry import Point, shape
from build_project import style

ROOT = Path(__file__).resolve().parents[2]
DESIGN = ROOT / 'design/final'
OUT = DESIGN / 'verification'
FINAL = ROOT / 'docs/final'

def table(headers, rows):
    def clean(x): return str(x).replace('|', '/').replace('\n', ' ')
    return '\n'.join(['| ' + ' | '.join(headers) + ' |', '| ' + ' | '.join(['---'] * len(headers)) + ' |'] +
                     ['| ' + ' | '.join(clean(x) for x in row) + ' |' for row in rows])

def box(ax, x, y, width, height, text, color='#eaf2f8', fontsize=10):
    ax.add_patch(FancyBboxPatch((x,y),width,height,boxstyle='round,pad=0.08',
                              edgecolor='#17324d',facecolor=color,linewidth=1))
    ax.text(x+width/2,y+height/2,text,ha='center',va='center',fontsize=fontsize,color='#17324d')

def arrow(ax, start, end):
    ax.annotate('', xy=end, xytext=start, arrowprops=dict(arrowstyle='->',color='#17324d',lw=1.3))

def node_models(data,z):
    unequal=data['topology'].get('zone_modes',{}).get(str(z))=='unequal'
    return ('PC252','PC250','PC252') if unequal else ('PC250','PC252','PC252')

def diagrams(data):
    fig, ax = plt.subplots(figsize=(15,10));ax.set_xlim(0,15);ax.set_ylim(0,10);ax.axis('off')
    ax.set_title('Подключение БС к головному узлу: две независимые MIMO-ветви',fontsize=16)
    for m,x in ((1,.4),(2,8.0)):
        ax.text(x+3.2,9.2,f'M{m}: отдельный RF-порт каждой LTE-несущей',ha='center',fontsize=12)
        for b,y in (('B3',6.7),('B1',4.4),('B7',2.1)):
            box(ax,x,y,2.1,1.35,f'LTE {b}\nМТС / О2 / О3\nпорт M{m}')
            box(ax,x+2.8,y,2.1,1.35,f'H-{b}-M{m}\nPC221, 4×4\nI1/I2/I3; I4')
            arrow(ax,(x+2.1,y+.7),(x+2.8,y+.7))
            box(ax,x+5.5,y,1.1,1.35,'O1–O4\nк TP\nZ1–Z4',fontsize=9)
            arrow(ax,(x+4.9,y+.7),(x+5.5,y+.7))
        ax.text(x+3.3,1.4,'I4 B1/B7: 50 Ом; нагрузки в перечне отсутствуют',ha='center',fontsize=9)
    box(ax,5.2,.1,4.5,.95,'Общий RF-выход GSM → S-DCS (PC250)\nO1 → H-B3-M1:I4; O2 → H-B3-M2:I4',fontsize=10)
    ax.text(7.5,8.8,'О2/О3: резерв LTE-входов. Общий GSM-выход двух TRX — допущение.',ha='center',fontsize=10)
    fig.savefig(OUT/'bs-connections.png',dpi=170,bbox_inches='tight');fig.savefig(OUT/'bs-connections.pdf',bbox_inches='tight');plt.close(fig)
    fig,ax=plt.subplots(figsize=(15,10));ax.set_xlim(0,15);ax.set_ylim(0,10);ax.axis('off')
    ax.set_title('Распределение одной MIMO-ветви Mj; вторая повторяет схему отдельно',fontsize=16)
    for z in range(1,5):
        y=9-(z-1)*2.25
        box(ax,.15,y-1.25,2.25,1.25,f'Выходы Oz трёх H\nB3/B1/B7 →\nTP-Z{z}-Mj, PC211')
        root,q1,q2=node_models(data,z)
        direct=[a['id'] for a in data['antennas'] if a['zone']==z and a['subgroup']==0]
        root_text=f'T{z}-Mj, 7/8\nS-Z{z}-Mj, {root}'
        root_text+=('\nO1/O2 → '+' / '.join(direct)+'\nO3/O4 → подгруппы') if direct else '\n2-way'
        box(ax,3.2,y-1.25,2.05,1.25,root_text,fontsize=8 if direct else 10)
        arrow(ax,(2.4,y-.65),(3.2,y-.65))
        for g,dy in ((1,0),(2,-1.0)):
            sy=y+dy-.15
            route=next(r for r in data['subfeeds'] if r['id']==f'F{z}-{g}')
            typ='джампер 3 м' if route['cable']=='jumper' else f'7/8: {route["length"]:.2f} м'
            box(ax,6.35,sy-.75,3.05,.75,f'F{z}-{g}-Mj: {typ}\nQ-Z{z}-{g}-Mj, {q1 if g==1 else q2}',fontsize=9)
            arrow(ax,(5.25,y-.65),(6.35,sy-.4))
            ids=[a['id'] for a in data['antennas'] if a['zone']==z and a['subgroup']==g]
            box(ax,10.3,sy-.75,4.4,.75,' / '.join(ids)+f'\nоконечные фидеры → джамперы → порт Pj',fontsize=9)
            arrow(ax,(9.4,sy-.4),(10.3,sy-.4))
    ax.text(7.5,-.35,'Z1–Z4 — группы распределения, не отдельные LTE-ячейки. M1 → P1; M2 → P2.',ha='center',fontsize=11)
    fig.savefig(OUT/'distribution-tree.png',dpi=170,bbox_inches='tight');fig.savefig(OUT/'distribution-tree.pdf',bbox_inches='tight');plt.close(fig)

def port_plan(data):
    rows=[]
    def add(source,target,link,state='Проект'):
        rows.append(dict(source=source,target=target,connection=link,status=state))
    for m in (1,2):
        for band in ('B3','B1','B7'):
            h=f'H-{band}-M{m}'
            for op in (1,2,3):
                add(f'LTE-{band}-О{op}:M{m}',f'{h}:I{op}','Джампер 3 м', 'МТС' if op==1 else 'Резерв, мощность не задана')
            if band!='B3':add('Нагрузка 50 Ом',f'{h}:I4','Не представлена в перечне','Неукомплектовано')
            for z in range(1,5):add(f'{h}:O{z}',f'TP-Z{z}-M{m}:{band}','Джампер 3 м')
        for z in range(1,5):
            add(f'TP-Z{z}-M{m}:COMMON',f'S-Z{z}-M{m}:IN',f'T{z}-M{m}, 7/8')
            for a in sorted([a for a in data['antennas'] if a['zone']==z and a['subgroup']==0],key=lambda a:a['id']):
                drop=next(r for r in data['drops'] if r['id']==a['id'])
                add(f'S-Z{z}-M{m}:O{a["leaf_output"]}',f'{a["id"]}:P{m}',f'{a["id"]}-M{m}, {drop["cable"]} + джампер 3 м')
            for g in (1,2):
                route=next(r for r in data['subfeeds'] if r['id']==f'F{z}-{g}')
                output=g+2 if node_models(data,z)[0]=='PC252' else g
                add(f'S-Z{z}-M{m}:O{output}',f'Q-Z{z}-{g}-M{m}:IN',f'F{z}-{g}-M{m}, '+('джампер 3 м' if route['cable']=='jumper' else '7/8'))
                for slot,a in enumerate([a for a in data['antennas'] if a['zone']==z and a['subgroup']==g],1):
                    drop=next(r for r in data['drops'] if r['id']==a['id'])
                    add(f'Q-Z{z}-{g}-M{m}:O{slot}',f'{a["id"]}:P{m}',f'{a["id"]}-M{m}, {drop["cable"]} + джампер 3 м')
    add('GSM:COMMON (2 TRX)', 'S-DCS:IN','Джампер 3 м','Общий выход БС принят условно')
    for m in (1,2):add(f'S-DCS:O{m}',f'H-B3-M{m}:I4','Джампер 3 м')
    assert len({r['target'] for r in rows})==len(rows), 'A destination port is connected twice'
    return rows

def device_layout(data):
    geom=json.loads((DESIGN/'cad-geometry.json').read_text())
    rows=[]
    def add(id,model,xy,basis):rows.append(dict(id=id,model=model,x_m=xy[0],y_m=xy[1],basis=basis))
    for m in (1,2):
        for band in ('B3','B1','B7'):add(f'H-{band}-M{m}','PC221',data['head'],'Головной узел; размеры и монтажная раскладка не заданы')
        for z in range(1,5):
            add(f'TP-Z{z}-M{m}','PC211',data['head'],'Головной узел')
            root,q1,q2=node_models(data,z)
            add(f'S-Z{z}-M{m}',root,data['trunks'][z-1]['points'][-1],'Зональное распределение по трассе')
            for g in (1,2):
                r=next(r for r in data['subfeeds'] if r['id']==f'F{z}-{g}')
                add(f'Q-Z{z}-{g}-M{m}',q1 if g==1 else q2,r['points'][-1],'Подгруппа на трассе; M1/M2 показаны одной геометрической точкой')
    add('S-DCS','PC250',data['head'],'Головной узел; общий GSM-выход условен')
    for a in data['antennas']:add(a['id'],'In06',(a['x'],a['y']),a['basis'])
    assert len(rows)==71 and len({r['id'] for r in rows})==71
    with (OUT/'device-locations.csv').open('w',newline='') as stream:
        writer=csv.DictWriter(stream,fieldnames=rows[0].keys(),lineterminator='\n');writer.writeheader();writer.writerows(rows)
    fig,ax=plt.subplots(figsize=(11,17))
    arch=[e['points'] for e in geom['lines'] if not any(v in e['layer'] for v in ('CABL','AREA'))]
    ax.add_collection(LineCollection(arch,colors='#999999',linewidths=.35))
    colors=['#bc3f46','#22846d','#ad701e','#774b9d']
    for z in range(1,5):
        color=colors[z-1]
        routes=[data['trunks'][z-1]]+[r for r in data['subfeeds'] if r['id'].startswith(f'F{z}-')]+[
            r for r in data['drops'] if next(a for a in data['antennas'] if a['id']==r['id'])['zone']==z]
        for r in routes:ax.plot(*zip(*r['points']),color=color,lw=1,alpha=.75)
        s=data['trunks'][z-1]['points'][-1];ax.scatter(*s,marker='s',s=80,color=color,zorder=5)
        joined=any(r['id']==f'F{z}-1' and r['cable']=='jumper' for r in data['subfeeds'])
        label=f'S-Z{z}' + (f' / Q-Z{z}-1' if joined else '')
        ax.annotate(label,s,xytext=(8,-14),textcoords='offset points',fontsize=7,color=color,
                    bbox=dict(facecolor='white',alpha=.9,edgecolor='none',pad=1))
        for g in (1,2):
            r=next(r for r in data['subfeeds'] if r['id']==f'F{z}-{g}')
            if r['cable']=='jumper':continue
            q=r['points'][-1];ax.scatter(*q,marker='D',s=50,color=color,zorder=5)
            ax.annotate(f'Q-Z{z}-{g}',q,xytext=(8,7),textcoords='offset points',fontsize=7,color=color,
                        bbox=dict(facecolor='white',alpha=.9,edgecolor='none',pad=1))
    for a in data['antennas']:
        color=colors[a['zone']-1];ax.scatter(a['x'],a['y'],s=28,color=color,zorder=6)
        ax.annotate(a['id'],(a['x'],a['y']),xytext=(3,3),textcoords='offset points',fontsize=6,color=color)
    ax.scatter(*data['head'],marker='*',s=150,color='#17324d',zorder=7)
    ax.annotate('ВК: H ×6; TP ×8; S-DCS ×1\nчасть 104.2; собственное решение',data['head'],xytext=(-80,45),
                textcoords='offset points',fontsize=7,color='#17324d',arrowprops=dict(arrowstyle='->'),
                bbox=dict(facecolor='white',alpha=.9,edgecolor='#17324d'))
    ax.set_xlim(-3,40);ax.set_ylim(-3,66);ax.set_aspect('equal');ax.grid(alpha=.12)
    ax.set_title('Первый этаж: один LTE-пул B3/B1/B7; 32 In06\nS — PC252; Q1 — PC250; Q2 — PC252; M1/M2 отдельно',fontsize=11)
    ax.set_xlabel('CAD X, м');ax.set_ylabel('CAD Y, м')
    fig.text(.5,.055,'S/Q: по два устройства, M1 и M2. Точные размеры и монтажная раскладка не заданы.\n'
             '● In06; ■ зональный делитель; ◆ делитель подгруппы; ★ головной узел. Проектные координаты, не натурная съёмка.',
             ha='center',fontsize=8)
    fig.savefig(OUT/'final-layout.png',dpi=190,bbox_inches='tight');fig.savefig(OUT/'final-layout.pdf',bbox_inches='tight');plt.close(fig)
    return rows

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    d=json.loads((DESIGN/'candidate.json').read_text())
    model=json.loads((DESIGN/'coverage/area-model.json').read_text())
    cov=json.loads((DESIGN/'coverage/coverage-report.json').read_text())
    fp=hashlib.sha256((DESIGN/'candidate.json').read_bytes()).hexdigest()
    assert fp==cov['design_fingerprint']
    groups=[shape(g['geometry']) for g in model['groups']]
    assert all(any(g.covers(Point(a['x'],a['y'])) for g in groups) for a in d['antennas'])
    assert all(next(shape(g['geometry']) for g in model['groups'] if a['room'] in g['rooms']).covers(Point(a['x'],a['y'])) for a in d['antennas'])
    assert any(shape(g['geometry']).covers(Point(d['head'])) and '104.2' in g['rooms'] for g in model['groups'])
    ports=port_plan(d)
    devices=device_layout(d)
    with (OUT/'port-plan.csv').open('w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=ports[0].keys(),lineterminator='\n');w.writeheader();w.writerows(ports)
    diagrams(d)
    with np.load(DESIGN/'coverage/validated-fields.npz') as field:
        xy=field['points'].copy();weights=field['weights'].copy();b3=field['rsrp_1800'].copy()
    # Same 1800 MHz ray geometry and reciprocal passive paths, an additional
    # 2-way DCS division, 0.5 dB assumed excess. Strongest individual path only.
    pilot=23-10*math.log10(200);split=10*math.log10(2)+.5
    dcs=[]
    for name,power,extra in [('D40_per_TRX',40,0),('D40_total_2TRX_equal',40-10*math.log10(2),0),
                             ('D30_per_TRX',30,0),('D40_RF_extra3dB',40,3)]:
        rx=b3+power-split-pilot-extra
        share=float(weights[rx>-85].sum()/2100*100)
        dcs.append(dict(scenario=name,bcch_source_dbm=power,extra_rf_db=extra,
                        reference_percent_2100=share,worst_rx_dbm=float(rx.min()),passed_95=share>=95))
        if name=='D40_per_TRX':
            np.savez_compressed(OUT/'dcs-field.npz',points=xy,weights=weights,rx_dbm=rx)
            fig,ax=plt.subplots(figsize=(9,14))
            plot=ax.scatter(*xy.T,c=rx,s=7,marker='s',cmap='RdYlGn',vmin=-100,vmax=-45)
            for g in groups:ax.plot(*g.exterior.xy,color='#333333',lw=.35)
            for a in d['antennas']:ax.plot(a['x'],a['y'],'k.',ms=3)
            ax.set_aspect('equal');ax.set_xlim(-2,39);ax.set_ylim(-2,66)
            ax.set_title(f'DCS 1800: условно 40 дБм на BCCH/TRX, n=2,5\nRx > −85: {share:.2f}% от 2100 м²; C/I не рассчитан',fontsize=10)
            ax.set_xlabel('CAD X, м');ax.set_ylabel('CAD Y, м');fig.colorbar(plot,ax=ax,label='Расчётная мощность BCCH, дБм')
            fig.savefig(OUT/'dcs-1800-conditional.png',dpi=170,bbox_inches='tight');plt.close(fig)
    fig,ax=plt.subplots(figsize=(9,14))
    ax.scatter(*xy.T,s=7,c='#2776af',marker='s')
    for g in groups:ax.plot(*g.exterior.xy,color='#333333',lw=.35)
    ax.set_aspect('equal');ax.set_xlim(-2,39);ax.set_ylim(-2,66)
    ax.set_title('Принадлежность локализованной площади проектному LTE-пулу\nОдна ячейка на диапазон; конкуренция с внешней сетью не задана',fontsize=10)
    ax.set_xlabel('CAD X, м');ax.set_ylabel('CAD Y, м')
    fig.savefig(OUT/'project-cell-assignment.png',dpi=170,bbox_inches='tight');plt.close(fig)
    routes=d['trunks']+d['subfeeds']+d['drops']
    conn={'1/2':{'male':0,'female':0},'7/8':{'male':0,'female':0}}
    for r in routes:
        if r['cable']=='jumper':continue
        if r['id'].startswith('A'):
            conn[r['cable']]['male']+=2;conn[r['cable']]['female']+=2
        else:conn[r['cable']]['male']+=4
    for cable in conn:assert sum(conn[cable].values())==4*sum(r['cable']==cable for r in routes)
    report=dict(date='2026-10-13',design_fingerprint=fp,antennas_inside_model=32,
                head_group='104.2',port_plan_connections=len(ports),device_locations=len(devices),
                manufactured_feeder_segments=2*sum(r['cable']!='jumper' for r in routes),
                local_jumper_count=d['local_jumper_count'],total_jumper_count=45+2*len(d['antennas'])+d['local_jumper_count'],
                connector_genders_required=conn,unused_input_loads_mts_only=16,unused_input_loads_three_lte_operators=4,
                dcs_scenarios=dcs,full_kpi_compliance=False)
    (OUT/'engineering-review.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    wb=load_workbook(ROOT/'equipment/bill-of-materials-draft.xlsx')
    for name in ('Портовый план','DCS сценарии','Комплектность','Размещение'):
        if name in wb:del wb[name]
    ws=wb.create_sheet('Портовый план');ws.append(['Источник','Назначение','Соединение','Статус'])
    for r in ports:ws.append(list(r.values()))
    style(ws,{'A':35,'B':35,'C':48,'D':50})
    ws=wb.create_sheet('Размещение');ws.append(['ID','Модель','X, м','Y, м','Основание / ограничение'])
    for r in devices:ws.append(list(r.values()))
    style(ws,{'A':28,'B':18,'C':22,'D':22,'E':110})
    ws=wb.create_sheet('DCS сценарии');ws.append(['Сценарий','Мощность BCCH дБм','Доп. потери дБ','Доля от 2100, %','Минимум дБм','Модель ≥95%'])
    for r in dcs:ws.append(list(r.values()))
    style(ws,{'A':32,'B':28,'C':24,'D':28,'E':24,'F':22})
    ws=wb.create_sheet('Комплектность');ws.append(['Элемент / требование','Количество / условие','Основание и ограничение'])
    for row in [
        ('Нагрузки входов, только МТС',16,'50 Ом; модели и допустимые мощности не предоставлены, в основную спецификацию не добавлены'),
        ('Нагрузки входов, три LTE-оператора',4,'Четвёртые входы гибридов B1/B7; отдельного резерва GSM О2/О3 нет'),
        ('Разъёмы 1/2', f'{conn["1/2"]["male"]} male + {conn["1/2"]["female"]} female','Это требуемые исполнения в рамках общей позиции перечня; доступность не подтверждена'),
        ('Разъёмы 7/8',f'{conn["7/8"]["male"]} male + {conn["7/8"]["female"]} female','Female на антенных концах для male/male джамперов; переходники не выдуманы'),
        ('In06 / БС','Интерфейс не задан','4.3-10F на In06 принят как условие стыковки, не характеристика из перечня'),
        ('Объединение двух GSM TRX','Условно общий RF-порт БС','При раздельных выходах нужно дополнительное решение; его модели нет в исходных данных'),
        ('DCS О2 / О3','Отдельных входов нет','Три LTE-оператора предусмотрены, полное трёхоператорское использование GSM не подтверждено'),
        ('Удалённое регулирование','Не выполнено пассивной схемой','Соответствующих модулей в перечне нет'),
        ('PIM −160дБн','Заявлено для пассивных компонентов в перечне','Условия измерения, допустимая суммарная мощность и изоляция не предоставлены'),
        ('Пожарное исполнение','XLSX: огнестойкое, безгалогеновое; PDF: PE','Расчёт затухания по PDF не подтверждает требуемую модификацию'),
    ]:ws.append(row)
    style(ws,{'A':38,'B':55,'C':100})
    for sheet in wb:
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
        sheet.print_title_rows='1:1';sheet.print_options.horizontalCentered=True
        sheet.print_area=sheet.dimensions
    wb.save(ROOT/'equipment/bill-of-materials-draft.xlsx')
    sections=['# Итоговая инженерная проверка\n\nДата: 13.10.2026. Данные задания, паспортные данные и собственные решения разделены. Полное соответствие ТЗ не установлено.\n',
              '## Исправления и проверенная согласованность\n\n'
              f'Местные связи на обеих ветвях: {d["local_jumper_count"]} заводских джамперов 3 м вместо нулевых фидеров. Добавлено 0,2 дБ оценочных потерь на затронутый путь. '
              f'Итого {report["manufactured_feeder_segments"]} изготовляемых фидера, {report["total_jumper_count"]} джамперов (12 резервных), {sum(conn["1/2"].values())} разъёмов 1/2 и {sum(conn["7/8"].values())} разъёма 7/8. '
              'Все 32 антенны лежат внутри восстановленных групп; ВК в 104.2. Это проверка модели, не разрешение на монтаж.\n',
              '## Неравномерное распределение мощности\n\n'
              'В каждой зоне PC252 питает две антенны непосредственно, две через PC250, четыре через второй PC252. '
              'Идеальные доли от зонального входа: 2×1/4 + 2×1/8 + 4×1/16 = 1. '
              'На стадиях деления относительно прежних восьми одинаковых путей непосредственные получают +3,51 дБ (включая исключение одной оценочной стадии потерь), '
              'парные сохраняют потери деления, четверные получают −3,01 дБ. Полные уровни включают отдельно пересчитанные кабели и местные соединения. Мощность пилота не изменена. Количества моделей делителей сохранены, их места в дереве изменены. '
              'Покрытие оценивается по всей площади, включая участки с меньшей мощностью и помещение 89 как непрошедшее.\n',
              '## Условный DCS Signal strength\n\n'
              'Данные задания: максимальная мощность GSM 40 дБм, RxLev >−85 дБм на 95%, C/I >20 дБ на 95%. '
              'Рабочая мощность BCCH, ARFCN и объединение двух TRX не заданы. Выбран сценарий 40 дБм на BCCH/TRX перед отдельным PC250; '
              'его потери 10log10(2)+0,5 = 3,5103 дБ. Половина GSM-мощности поступает в каждое дерево; мощности двух антенных портов не суммируются. '
              'Применены те же 1800 МГц, высоты, n=2,5, препятствия и веса, что в LTE.\n\n'
              'Rx_DCS = RSRP_B3,P1 + P_BCCH − (10log10(2)+0,5) − P_CRS_RE,B3.\n\n'
              'Это прогноз мощности BCCH, а не измеренный RxLev. Максимальные 40 дБм не объявляются рабочей настройкой. '
              'Если 40 дБм относятся к сумме двух равномощных TRX, мощность одного равна 36,99 дБм.\n\n'+
              table(['Сценарий','BCCH, дБм','Доля Rx >−85 от 2100, %'],[(r['scenario'],f'{r["bcch_source_dbm"]:.2f}',f'{r["reference_percent_2100"]:.2f}') for r in dcs])+ '\n',
              '## RSRQ, SINR, C/I и UL: что следует из доступных данных\n\n'
              'RSRQ_dB = 10log10(N_RB) + RSRP_dBm − RSSI_dBm, при согласованной полосе и определении измерения. '
              'RSRQ >−12 требует RSSI < RSRP + 10log10(N_RB)+12. При RSRP=−90 пределы RSSI для B3/B1/B7 составляют −58,00/−59,25/−61,01 дБм. '
              'RSSI и режим измерения не заданы; это пределы, а не расчётные карты или выполненный KPI.\n\n'
              'SINR = S/(I+N). При мощности S=−90 дБм на согласованном RE пороги 17/7/1 дБ требуют I+N ≤−107/−97/−91 дБм. '
              'Эти границы не определяют доли 40/80/95%: нет внешних помех, загрузки и шумовой модели. '
              'По DCS C/I >20 требует I < C−20; при C=−85 это I <−105 дБм. ARFCN и co-channel-план неизвестны.\n\n'
              'UL: P_UE + G_UE + G_In06 − L_propagation,UL − L_walls − L_RF,UL ≥ S_BS + margin. '
              'Пассивные потери учитываются на частоте UL; из DL-таблицы нельзя автоматически получить UL-чувствительность. '
              'Мощность UE, NF, S_BS, настройки приёмника и запас не заданы, поэтому численный UL-KPI не заявляется.\n',
              '## Секторизация и Best Server\n\n'
              'На каждом диапазоне проект содержит одну LTE-ячейку МТС. Z1–Z4 — распределительные группы. '
              'Карта принадлежности к проектной ячейке не является полноценной картой Best Server с внешней macro-сетью. '
              'Антенна с максимальным отдельным путём не считается новой ячейкой. EARFCN/PCI и конкурирующая сеть не заданы.\n',
              '## Комплектность и стыковка\n\n'
              'При работе только МТС нужно закончить 16 свободных входов гибридов; при подключении трёх LTE-операторов остаются четыре свободных входа B1/B7. '
              'Согласованные нагрузки 50 Ом отсутствуют в перечне, их мощность неизвестна. В спецификацию не добавлялись выдуманные позиции. '
              'Доступность исполнений разъёмов male/female, интерфейса In06 и переходов от БС не подтверждена. '
              f'Потребность: 1/2 — {conn["1/2"]["male"]} male+{conn["1/2"]["female"]} female; 7/8 — {conn["7/8"]["male"]} male+{conn["7/8"]["female"]} female. Антенные female-окончания нужны для male/male джамперов. '
              'Для пассивных устройств перечень заявляет −160дБн, но не задаёт условия PIM, изоляцию и допустимую суммарную мощность. '
              'Ни эти свойства, ни регулирование удалённых модулей, ни пожарное исполнение не объявляются подтверждёнными.\n',
              '## Итоговое решение\n\n'
              'Сохранён один общий трёхдиапазонный LTE-пул и 32 In06. Он обслуживает 13 UE по 8 Мбит/с в принятой модели ёмкости и проходит номинальный прогноз RSRP. '
              'Запаса на 14-го UE без CA нет. Проверки +3 дБ и закрытых проёмов приведены в отчёте покрытия; более сильное затухание и иная нормировка пилота остаются отдельными ограничениями. Полное соответствие задания не доказано. '
              'Второй сектор сам по себе не исправляет стены, UL, нормировку CRS и комплектацию, поэтому не добавлен без отдельного обоснования. '
              'Недоступные исходные данные не заменены благоприятными выдуманными настройками.\n',
              '## Производные файлы\n\n'
              '`design/final/verification/`: схемы подключения БС и дерева, портовый план CSV, карта DCS, карта принадлежности проектной ячейке, численный отчёт. '
              'В XLSX добавлены листы «Портовый план», «DCS сценарии», «Комплектность».\n']
    (FINAL/'engineering-review.md').write_text('\n'.join(sections))
    print(json.dumps(report,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
