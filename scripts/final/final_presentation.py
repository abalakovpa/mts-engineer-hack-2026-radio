#!/usr/bin/env python3
"""Build six editable slides using the supplied MTS template and render the PPTX."""
from copy import deepcopy
from pathlib import Path
import hashlib
import json
import math
import shutil
import subprocess
import tempfile

import fitz
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'presentation'
TEMPLATE=OUT/'source-template.pptx'
TEMPLATE_SHA='83f1fd8e22427bed6954db9cccb7145e5662e14aa89208947e42e5fab1598a28'
RED='FF0032'
INK='414141'
MUTED='666666'
TEXT='MTS Text'
WIDE='MTS Wide Medium'
NAV=['Автор','Ёмкость','Схема DAS','Покрытие','Оборудование','Итоги']
pdfmetrics.registerFont(TTFont('RadioMeasure','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))


def n(value,places=2):
    return f'{value:.{places}f}'.replace('.',',')


def wrap(value,size,width):
    """Conservative line breaks; the actual template fonts are checked after render."""
    lines=[]
    for paragraph in value.split('\n'):
        current=''
        for word in paragraph.split():
            candidate=(current+' '+word).strip()
            if current and pdfmetrics.stringWidth(candidate,'RadioMeasure',size)>width*72:
                lines.append(current);current=word
            else:current=candidate
        lines.append(current)
    return lines


class Deck:
    def __init__(self):
        assert hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()==TEMPLATE_SHA
        self.prs=Presentation(TEMPLATE)
        assert len(self.prs.slides)==8
        assert self.prs.slide_width==Inches(10) and self.prs.slide_height==Inches(5.625)
        # Preserve the original gradient and brand assets. Unused sample slides are removed.
        gradient=deepcopy(self.prs.slides[2]._element.cSld.bg)
        for item in list(self.prs.slides._sldIdLst)[6:]:
            self.prs.part.drop_rel(item.rId);self.prs.slides._sldIdLst.remove(item)
        for index,slide in enumerate(self.prs.slides):
            for rel in list(slide.part.rels.values()):
                if rel.reltype==RT.NOTES_SLIDE:
                    slide.part.drop_rel(rel.rId)
            if index==1:
                old=slide._element.cSld.bg
                if old is not None:slide._element.cSld.remove(old)
                slide._element.cSld.insert(0,deepcopy(gradient))
            for shape in list(slide.shapes):
                is_brand=shape.shape_type==13 and (
                    shape.left>Inches(9) or (index==0 and shape.top<Inches(.3) and shape.left<Inches(.3)))
                if not is_brand:shape._element.getparent().remove(shape._element)
        self.slide=None;self.index=0

    def rect(self,x,y,w,h,color='FFFFFF',rounded=True,line=None):
        shape=self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                           Inches(x),Inches(y),Inches(w),Inches(h))
        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(color)
        if rounded:shape.adjustments[0]=.1
        if line:
            shape.line.color.rgb=RGBColor.from_string(line);shape.line.width=Pt(.75)
        else:shape.line.fill.background()
        shape._element.spPr.append(OxmlElement('a:effectLst'))
        return shape

    def text(self,x,y,w,h,value,size=14,heading=False,color=INK,align=PP_ALIGN.LEFT):
        lines=wrap(value,size,w-.04)
        leading=size*1.18
        assert len(lines)*leading<=h*72+.5,(self.index,value,len(lines)*leading,h*72)
        shape=self.slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        frame=shape.text_frame
        frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
        frame.word_wrap=False;frame.auto_size=MSO_AUTO_SIZE.NONE
        for i,line in enumerate(lines):
            p=frame.paragraphs[0] if i==0 else frame.add_paragraph()
            p.text=line;p.font.name=WIDE if heading else TEXT;p.font.size=Pt(size)
            p.font.color.rgb=RGBColor.from_string(color);p.alignment=align
            p.space_before=p.space_after=Pt(0);p.line_spacing=Pt(leading)
        return shape

    def image(self,path,x,y,w,h):
        with Image.open(path) as image:iw,ih=image.size
        factor=min(w/iw,h/ih);dw,dh=iw*factor,ih*factor
        self.slide.shapes.add_picture(str(path),Inches(x+(w-dw)/2),Inches(y+(h-dh)/2),
                                      width=Inches(dw),height=Inches(dh))

    def line(self,x1,y1,x2,y2,color=RED):
        shape=self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
        shape.line.color.rgb=RGBColor.from_string(color);shape.line.width=Pt(1.3)

    def new(self,title):
        self.slide=self.prs.slides[self.index];self.index+=1
        if self.index>1:
            for i,label in enumerate(NAV):
                x=.2+i*1.43;active=i==self.index-1
                self.rect(x,.2,1.36,.28,RED if active else 'FFFFFF')
                self.text(x,.255,1.36,.18,label,8.4,color='FFFFFF' if active else MUTED,align=PP_ALIGN.CENTER)
            self.text(.2,.78,9.3,1.0,title,28,True)
        self.text(8.88,5.42,.87,.13,f'{self.index} / 6',7,color=MUTED,align=PP_ALIGN.RIGHT)

    def card(self,x,y,w,h,label,value,detail):
        self.rect(x,y,w,h)
        self.text(x+.17,y+.14,w-.34,.25,label,10,color=RED)
        self.text(x+.17,y+.52,w-.34,.42,value,17.2 if len(value)>14 else 20,True)
        self.text(x+.17,y+1.02,w-.34,h-1.1,detail,11.2)

    def finish(self):
        self.prs.core_properties.title='Радиосеть первого этажа ТРК «Эльбрус»'
        self.prs.core_properties.author='Павел Абалаков'
        self.prs.core_properties.subject='МТС Engineer Hack 2026 · Радио'
        self.prs.core_properties.comments=''
        self.prs.save(OUT/'defense.pptx')


def render_and_check():
    # Render the actual editable PPTX, including the supplied embedded fonts.
    with tempfile.TemporaryDirectory(prefix='radio-ppt-render-',dir='/tmp') as temp:
        result=subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',temp,str(OUT/'defense.pptx')],
                              text=True,capture_output=True,timeout=90)
        if result.returncode:raise RuntimeError(result.stdout+result.stderr)
        source=Path(temp)/'defense.pdf'
        assert source.exists(),result.stdout+result.stderr
        shutil.copyfile(source,OUT/'defense-preview.pdf')
    with fitz.open(OUT/'defense-preview.pdf') as doc:
        assert len(doc)==6
        all_text='\n'.join(page.get_text() for page in doc)
        for forbidden in ('GitHub','github.com','[заполнить]','Навигация','Тезис слайда','несоответствие','Оформление v2'):
            assert forbidden not in all_text,forbidden
        for page in doc:
            for block in page.get_text('dict')['blocks']:
                for line in block.get('lines',[]):
                    x0,y0,x1,y1=line['bbox']
                    assert x0>=0 and y0>=0 and x1<=page.rect.width+.5 and y1<=page.rect.height+.5,('Rendered slide bounds',page.number+1,line)
        assert '13' in doc[1].get_text() and '104' in doc[1].get_text()
        assert '32' in doc[2].get_text() and '96,37' in doc[3].get_text()
        assert '1366' in doc[4].get_text() and '179' in doc[4].get_text()
        fonts={doc.extract_font(item[0])[0] for page in doc for item in page.get_fonts()}
        assert any('MTS' in font for font in fonts),fonts
        preview=Image.new('RGB',(1600,1410),'#E9E9E9');draw=ImageDraw.Draw(preview)
        for i,page in enumerate(doc):
            pix=page.get_pixmap(matrix=fitz.Matrix(1.6,1.6))
            image=Image.frombytes('RGB',(pix.width,pix.height),pix.samples);image.thumbnail((780,440))
            x=(i%2)*800;y=(i//2)*470
            preview.paste(image,(x+10,y+10));draw.text((x+15,y+451),f'{i+1} / 6',fill='#414141')
        preview.save(OUT/'slides-preview.png',optimize=True)
        return sorted(fonts)


def main():
    data=json.loads((ROOT/'design/final/candidate.json').read_text())
    report=json.loads((ROOT/'design/final/coverage/coverage-report.json').read_text())
    profile=json.loads((OUT/'profile-input.json').read_text())
    quality=json.loads((ROOT/'design/final/quality/quality-report.json').read_text())
    assert quality['design_fingerprint']==report['design_fingerprint']
    lte=next(r for r in quality['lte_results'] if r['scenario']=='L2' and r['frequency_mhz']==2600)
    dcs=next(r for r in quality['dcs_results'] if r['bcch_dbm']==40 and r['interference_dbm_200khz']==-70)
    values={(r['scenario'],r['frequency_mhz']):r['reference_percent_2100'] for r in report['scenarios']}
    total=data['cable78_installed']+data['cable12_installed']
    buy78=math.ceil(data['cable78_installed']*1.1);buy12=math.ceil(data['cable12_installed']*1.1)
    n3=values['P1_n3',2600];closed=values['P1_n3_closed_gaps',2600]
    scripts=[
        ('Об авторе и задаче',25,profile['self_introduction']+' В этом кейсе я проектирую радиосеть первого этажа ТРК «Эльбрус»: LTE в трёх диапазонах и DCS-1800.'),
        ('Ёмкость',40,
         'В онлайн-модели одновременно активными считались 80 процентов прогнозных пользователей. Новые вводные задают две последовательные доли: RRC Connected и наличие данных в буфере. '
         'Из 360 прогнозных пользователей получаются 126 подключений и 12,6 пользователя с данными. Округляю вверх до 13: необходимы 104 мегабита в секунду. '
         'При принятой эффективности три несущие обслуживают шесть, четыре и три устройства без агрегации. Выбран один общий LTE-сектор. '
         'Последовательность числа секторов: 25 в онлайн-этапе, два при новых долях и прежнем коэффициенте, один при общем пуле с коэффициентом единица. Резерва для четырнадцатого устройства без CA нет.'),
        ('Архитектура DAS',55,
         'На плане размещены 32 двухпортовые антенны In06. Четыре группы распределяют сигнал одного LTE-пула на каждом диапазоне. '
         'До двух портов каждой антенны идут независимые RF-деревья M1 и M2. На каждой ветви три гибрида PC221 распределяют диапазоны на четыре направления. '
         'Триплексеры объединяют B3, B1 и B7 в зональный фидер. В группе две антенны подключены напрямую, две через двухвыходный делитель и четыре через четырёхвыходный. '
         'Доли мощности — четверть, восьмая и шестнадцатая от зонального входа до дополнительных потерь. '
         'Головной узел предусмотрен в проектной зоне 104.2. Высоты антенны и терминала приняты семь и полтора метра. GSM использует отдельный делитель и общий выход двух TRX как условие подключения БС.'),
        ('Покрытие и качество',55,
         'RF-бюджет и карты учитывают кабели, высоты и пересечения стен. Пилот 23 дБм принят как сумма CRS-поднесущих одного порта, модель P1. '
         f'На B7 при n=3 покрыто {n(n3)} процента площади. '
         'Для SINR и RSRQ заданы NF семь децибел, загрузка 80 процентов и EPRE данных на три децибела ниже CRS. '
         'В L2 внешняя помеха равна минус 115 дБм на 15 килогерц. '
         f'На B7 SINR не ниже 17 дБ достигается на {n(lte["sinr_ge17_percent"])} процента, RSRQ выше минус 12 — на {n(lte["rsrq_gt_minus12_percent"])} процента. '
         'При равной EPRE критерий RSRQ не проходит. '
         f'DCS при BCCH 40 и помехе минус 70 дБм даёт C/I выше 20 на {n(dcs["ci_gt20_percent"])} процента. '
         'Это сценарные условия, а не параметры измеренной сети. P2 сохраняет непроходящий RSRP.'),
        ('Оборудование',35,
         'Количество оборудования получено из той же схемы и трасс. Предусмотрены 32 антенны, шесть гибридов, восемь триплексеров, девять двухвыходных и шестнадцать четырёхвыходных делителей. '
         f'Установлено {n(total,0)} метра коаксиала. С десятипроцентным резервом нужны {buy78} метров семь восьмых и {buy12} метров одна вторая. '
         'Для соединений учтены 168 разъёмов и 113 джамперов: 101 рабочий и 12 резервных. '
         'Свободные входы заканчиваются нагрузками 50 Ом. Их модели и допустимые мощности выбираются при комплектации. Огнестойкое безгалогеновое исполнение кабелей является условием поставки.'),
        ('Итоги',30,
         'Результат опирается на три решения. Последовательные доли RRC и buffer определяют спрос и один общий LTE-сектор. '
         '32 антенны и две MIMO-ветви обеспечивают распределение сигнала по этажу. Схема два плюс два плюс четыре задаёт разные доли мощности внутри групп. '
         'Расчётные данные согласованы между PDF, схемами и ведомостью оборудования. Покрытие относится к модели P1, качество рассчитано при заданных помехах и мощности данных. '
         'Отдельной проверки требуют UL и параметры поставочных исполнений. Архитектура пассивная: управляемые удалённые модули в её состав не входят.')]
    assert sum(row[1] for row in scripts)==240
    deck=Deck()
    deck.new('')
    deck.text(.25,1.78,6.65,1.1,'Радиосеть первого этажа\nТРК «Эльбрус»',29,True)
    deck.rect(.25,3.06,4.55,.46,line=RED)
    deck.text(.4,3.15,4.25,.3,profile['name'],18,True,color=RED)
    deck.text(.25,3.77,6.35,.85,'УрТИСИ · III курс · 11.03.02\nИнфокоммуникационные технологии\nи системы связи',14)
    deck.text(.25,4.72,6.3,.63,profile['profile_facts']+'\nАрхитектура, задачи и коммуникация\nИнтересы: product management и управление проектами',11.2)
    deck.rect(7.05,1.65,2.6,3.45)
    photos=[path for path in (OUT/'assets/portrait.jpg',OUT/'assets/portrait.png') if path.exists()]
    assert len(photos)<=1, 'Keep a single portrait asset'
    photo=photos[0] if photos else OUT/'assets/portrait.jpg'
    if photo.exists():deck.image(photo,7.15,1.75,2.4,3.25)
    else:
        deck.text(7.25,2.62,2.2,.85,'ПА',46,True,color=RED,align=PP_ALIGN.CENTER)
        deck.text(7.2,3.64,2.3,.6,'Павел\nАбалаков',17,True,align=PP_ALIGN.CENTER)
    deck.new('Новые доли RRC и buffer:\n13 пользователей вместо 288')
    for x,label,value,detail in [(.2,'ОДНОВРЕМЕННАЯ НАГРУЗКА','13 UE','6 + 4 + 3 по несущим'),
                                (3.43,'ПОЛЕЗНЫЙ DOWNLINK','104 Мбит/с','13 устройств по 8 Мбит/с'),
                                (6.66,'ОБЩАЯ ЁМКОСТЬ','1 LTE-сектор','Три несущие B3 / B1 / B7')]:
        deck.card(x,2.06,3.14,1.48,label,value,detail)
    deck.rect(.2,3.7,5.15,1.48);deck.rect(5.5,3.7,4.3,1.48)
    deck.text(.38,3.9,4.8,.38,'360 × 35% × 10% = 12,6 → 13 UE',16,True)
    deck.text(.38,4.45,4.8,.57,'Секторы: 25 → 2 (k=1,10) → 1 (k(1)=1)\nη=3,1505922; u=0,8.',11.8)
    deck.text(5.7,3.89,3.9,.24,'Полезная ёмкость трёх несущих',11,color=MUTED)
    for i,(band,bw,c,ue) in enumerate([('B3',20,50.4094752,6),('B1',15,37.8071064,4),('B7',10,25.2047376,3)]):
        deck.text(5.7,4.24+i*.25,3.9,.23,f'{band} · {bw} МГц · {n(c)} Мбит/с · {ue} UE',11.4)
    deck.text(.25,5.33,8.5,.18,'Без CA резерв для 14-го UE отсутствует; прежнее правило k=1,10 даёт два сектора.',9.1,color=MUTED)
    deck.new('Один LTE-пул — четыре группы\nи два независимых RF-дерева')
    deck.rect(.2,1.99,3.45,3.24);deck.rect(3.8,1.99,6.0,3.24)
    deck.image(ROOT/'design/final/verification/final-layout.png',.32,2.08,3.21,2.87)
    deck.text(.42,5.0,3.0,.19,'План первого этажа · 32 In06',10,color=MUTED,align=PP_ALIGN.CENTER)
    deck.text(4.0,2.14,5.6,.33,'На каждой MIMO-ветви',17,True)
    deck.text(4.0,2.64,5.6,.35,'PC221 ×3 → PC211 ×4 → группы Z1–Z4',15)
    for i in range(4):
        x=4.0+i*1.42;deck.rect(x,3.16,1.28,.57,'FFF1F5',line=RED)
        deck.text(x,3.28,1.28,.26,f'Z{i+1}: 8 In06',12,align=PP_ALIGN.CENTER)
    deck.text(4.0,3.98,5.6,.33,'В группе: 2 напрямую + 2 через PC250 + 4 через PC252',12.8)
    deck.text(4.0,4.59,5.6,.36,'Доли мощности: 1/4 · 1/8 · 1/16\nM1 и M2 подключены к разным портам антенн.',11)
    deck.text(.25,5.33,8.5,.18,'Головной узел — проектная зона 104.2. Высоты антенны / UE: 7 / 1,5 м.',9.1,color=MUTED)
    deck.new('RSRP и качество DL:\nрезультаты условных сценариев')
    deck.rect(.2,1.99,3.45,3.24);deck.rect(3.8,1.99,6.0,3.24)
    deck.image(ROOT/'design/final/coverage/rsrp-2600-n3.png',.33,2.09,3.19,2.92)
    deck.text(.4,5.01,3.05,.18,'B7 · n=3 · карта с шагом 0,25 м',9.5,color=MUTED,align=PP_ALIGN.CENTER)
    deck.text(4.0,2.13,5.6,.25,'RSRP > −90 дБм · доля от 2100 м²',14,True)
    xs=[4.0,6.55,7.55,8.6]
    for x,w,label in zip(xs,[2.4,.8,.8,.9],['Расчётный режим','B3','B1','B7']):deck.text(x,2.67,w,.23,label,11,color=MUTED)
    scenarios=[('P1_n3','n=3'),('P1_n3_closed_gaps','n=3, закрытые проёмы'),('P2_carrier23','P2: 23 дБм на несущую')]
    for i,(key,label) in enumerate(scenarios):
        y=3.03+i*.3;color=RED if key.startswith('P2') else INK
        deck.text(xs[0],y,2.5,.24,label,10.6,color=color)
        for x,f in zip(xs[1:],(1800,2100,2600)):deck.text(x,y,.95,.24,n(values[key,f])+'%',11,color=color)
    deck.text(4.0,4.05,5.55,.42,f'L2 / B7: SINR ≥17 — {n(lte["sinr_ge17_percent"])}%\nRSRQ >−12 — {n(lte["rsrq_gt_minus12_percent"])}%',11.5)
    deck.text(4.0,4.68,5.55,.25,f'DCS: C/I >20 — {n(dcs["ci_gt20_percent"])}% при I=−70 дБм',10.5)
    deck.text(.25,5.33,8.5,.18,'L2: P1, n=2,5; NF=7 дБ, β=0,5, ρ=0,8; I=−115 дБм/15 кГц. DCS: I в полосе 200 кГц.',8.8,color=MUTED)
    deck.new('Оборудование и трассы\nследуют одной схеме')
    cards=[('АНТЕННЫ','32 In06','Два RF-порта на антенну\n64 антенных порта'),
           ('ГОЛОВНОЙ УЗЕЛ','6 PC221 · 8 PC211','Объединение операторов\nи диапазонов B3 / B1 / B7'),
           ('РАСПРЕДЕЛИТЕЛИ','9 PC250 · 16 PC252','Парные и четверные группы\nСхема распределения 2+2+4'),
           ('КОАКСИАЛ',f'{n(total,0)} м',f'С резервом: {buy78} м 7/8\nи {buy12} м 1/2'),
           ('СОЕДИНЕНИЯ','168 разъёмов','7/8 — 100 шт.; 1/2 — 68 шт.\nИсполнения male / female'),
           ('ДЖАМПЕРЫ','113 × 3 м','101 рабочий + 12 резервных\nВключая четыре местных')]
    for i,(label,value,detail) in enumerate(cards):deck.card(.2+(i%3)*3.23,1.99+(i//3)*1.65,3.14,1.53,label,value,detail)
    deck.text(.25,5.33,8.5,.18,'Условия комплектации: нагрузки 50 Ом, интерфейсы RF-портов и пожарное исполнение кабелей.',9.1,color=MUTED)
    deck.new('Три решения, которые\nопределяют результат')
    for i,(label,value,detail) in enumerate([('МОДЕЛЬ СПРОСА','1 LTE-сектор','13 пользователей с данными\n104 Мбит/с полезного DL'),
                                          ('РАЗМЕЩЕНИЕ','32 антенны',f'B7 при n=3: {n(n3)}%\nДве независимые MIMO-ветви'),
                                          ('ДЕЛЕНИЕ МОЩНОСТИ','2 + 2 + 4','Доли 1/4, 1/8 и 1/16\nв каждой группе')]):
        deck.card(.2+i*3.23,2.05,3.14,1.64,label,value,detail)
    deck.rect(.2,3.91,9.6,1.28)
    deck.text(.42,4.09,9.15,.33,'Ёмкость → архитектура → RF-бюджет → покрытие → спецификация',17,True)
    deck.text(.42,4.64,9.15,.4,'SINR, RSRQ и C/I рассчитаны при явных параметрах помех и EPRE. UL и параметры поставки\nпроверяются отдельно; управляемые удалённые модули в состав пассивной сети не входят.',11)
    deck.finish()
    fonts=render_and_check()
    word_count=sum(len(speech.split()) for _,_,speech in scripts)
    text=['# Выступление на четыре минуты\n\nПлановый тайминг — 240 секунд. '
          f'Текст содержит {word_count} слов; фактическую длительность проверяет докладчик.\n']
    for i,(title,seconds,speech) in enumerate(scripts,1):text.append(f'## Слайд {i}. {title} — {seconds} секунд\n\n{speech}\n')
    text.append('## Репетиция\n\nПервый проход — объяснить модель по слайдам. Второй — выступить с таймером. '
                'Повторить определения RRC, buffer users, CRS RE и различие между антенной группой и сектором. '
                'Для вопросов подготовить один полный RF-путь и ведомость портов.\n')
    (OUT/'speaker-script.md').write_text('\n'.join(text))
    pending=[key for key in ('role','contribution','interests') if not profile[key]]
    (OUT/'presentation-status.json').write_text(json.dumps(dict(slides=6,planned_seconds=240,word_count=word_count,
            pending_profile_fields=pending,official_template_available=True,photo_available=photo.exists(),
            template_sha256=hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),rendered_from_pptx=True,
            quality_report_sha256=hashlib.sha256((ROOT/'design/final/quality/quality-report.json').read_bytes()).hexdigest(),
            rendered_fonts=fonts,speaker_notes_included=False,personal_rehearsal_done=False),ensure_ascii=False,indent=2)+'\n')
    print(f'Presentation: six slides, {word_count} words, official template, actual PPTX render checked; photo={photo.exists()}')

if __name__=='__main__':main()
