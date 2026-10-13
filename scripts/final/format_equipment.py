#!/usr/bin/env python3
"""Apply the report theme to all equipment sheets without changing cell values."""
from pathlib import Path
from openpyxl import load_workbook
from build_project import style

ROOT=Path(__file__).resolve().parents[2]

def format_equipment():
    path=ROOT/'equipment/bill-of-materials.xlsx'
    workbook=load_workbook(path)
    for sheet in workbook:
        widths={sheet.cell(1,i).column_letter:
                sheet.column_dimensions[sheet.cell(1,i).column_letter].width
                for i in range(1,sheet.max_column+1)}
        style(sheet,widths)
    workbook.save(path)
    print(f'Equipment theme applied to {len(workbook.sheetnames)} sheets.')

if __name__=='__main__':format_equipment()
